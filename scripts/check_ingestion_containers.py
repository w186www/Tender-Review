"""Mandatory real Docker/S3/ClamAV/OCR/HTTP gate; owns and removes only its test resources."""

import io
import os
import shutil
import socket
import subprocess
import sys
import time
from pathlib import Path
from uuid import uuid4

import httpx
from docx import Document as WordDocument
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session

from bidding_compliance.core.config import Settings
from bidding_compliance.ingestion.policy import load_policy
from bidding_compliance.ingestion.ports import IngestionError
from bidding_compliance.ingestion.sandbox import DockerParser
from bidding_compliance.ingestion.scanner import ClamAVScanner
from bidding_compliance.ingestion.storage import S3ObjectStore
from bidding_compliance.persistence.database import make_engine, migrate
from bidding_compliance.persistence.models import Tenant, User
from bidding_compliance.services.identity import PASSWORDS
from bidding_compliance.services.ingestion_worker import execute

ROOT = Path(__file__).resolve().parents[1]
DOCKER = shutil.which("docker")


def docker(arguments: list[str], environment: dict[str, str] | None = None) -> str:
    if DOCKER is None:
        raise RuntimeError("Docker is required for this gate")
    result = subprocess.run(  # noqa: S603 - fixed executable, controlled arguments
        [DOCKER, *arguments],
        env=environment,
        capture_output=True,
        text=True,
        check=True,
        timeout=60,
    )
    return result.stdout.strip()


def pdf() -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(40, 700, "Synthetic tender qualification and contract requirement.")
    document.showPage()
    document.save()
    return buffer.getvalue()


def docx() -> bytes:
    buffer = io.BytesIO()
    document = WordDocument()
    document.add_heading("Synthetic qualification", 1)
    document.add_paragraph("项目名称：合成工程。政府采购工程项目。保留原文、页码或稳定定位。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Synthetic requirement"
    table.cell(0, 1).text = "Test only"
    document.save(buffer)
    return buffer.getvalue()


def image() -> bytes:
    result = Image.new("RGB", (1800, 240), "white")
    fonts = [
        ROOT / ".verification/not-used.ttf",
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    font = next(
        (ImageFont.truetype(str(p), 48) for p in fonts if p.exists()),
        ImageFont.load_default(size=48),
    )
    ImageDraw.Draw(result).text(
        (40, 70), "Synthetic tender qualification requirements", font=font, fill="black"
    )
    buffer = io.BytesIO()
    result.save(buffer, format="PNG")
    return buffer.getvalue()


def scanned_pdf() -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawImage(ImageReader(io.BytesIO(image())), 40, 700, width=500, height=67)
    document.showPage()
    document.save()
    return buffer.getvalue()


def wait_for_scanner(scanner: ClamAVScanner, timeout_seconds: float = 240) -> None:
    """Wait for startup signature updates; never accept an unavailable/stale scanner."""
    deadline = time.monotonic() + timeout_seconds
    while True:
        try:
            scanner.scan(b"Synthetic clean scanner test.")
            return
        except IngestionError as exc:
            if exc.code not in {"SCANNER_UNAVAILABLE", "SCANNER_STALE"}:
                raise
            if time.monotonic() >= deadline:
                raise RuntimeError(f"ClamAV readiness failed: {exc.code}") from exc
            time.sleep(min(2, max(0, deadline - time.monotonic())))


def main() -> None:
    parser_image = os.environ.get("PARSER_VERIFY_IMAGE")
    if not parser_image:
        raise RuntimeError("Build the reviewed image and set PARSER_VERIFY_IMAGE to its image ID")
    parser = DockerParser(parser_image)
    names: list[str] = []
    server = None
    engine = None
    folder = ROOT / ".verification" / ("ingestion-" + uuid4().hex)
    folder.mkdir(parents=True)
    log = (folder / "server.log").open("wb")
    unique = uuid4().hex[:10]
    environment = os.environ.copy()
    environment.update(MINIO_ROOT_USER="synthetic-" + unique, MINIO_ROOT_PASSWORD=str(uuid4()))
    try:
        minio, clam = "tender-minio-" + unique, "tender-clamav-" + unique
        # Always resolve the reviewed local build to an immutable image ID.
        minio_image = docker(
            [
                "image",
                "inspect",
                os.environ.get("MINIO_VERIFY_IMAGE", "tender-minio-verification:task5"),
                "--format",
                "{{.Id}}",
            ]
        )
        clam_image = docker(["image", "inspect", "clamav/clamav:stable", "--format", "{{.Id}}"])
        for name, args in (
            (
                minio,
                [
                    "-p",
                    "127.0.0.1::9000",
                    "-e",
                    "MINIO_ROOT_USER",
                    "-e",
                    "MINIO_ROOT_PASSWORD",
                    minio_image,
                    "server",
                    "/data",
                ],
            ),
            (clam, ["-p", "127.0.0.1::3310", "--memory=3g", clam_image]),
        ):
            docker(
                [
                    "run",
                    "--detach",
                    "--rm",
                    "--name",
                    name,
                    "--label",
                    "tender-verification=task5",
                    *args,
                ],
                environment,
            )
            names.append(name)
        storage_port = int(docker(["port", minio, "9000"]).rsplit(":", 1)[1])
        scanner_port = int(docker(["port", clam, "3310"]).rsplit(":", 1)[1])
        password, key = str(uuid4()), str(uuid4()) + str(uuid4())
        settings = Settings(
            _env_file=None,
            app_env="test",
            database_url="sqlite:///" + (folder / "data.db").as_posix(),
            auth_signing_key=key,
            object_storage_mode="s3",
            object_storage_endpoint=f"http://127.0.0.1:{storage_port}",
            object_storage_bucket="synthetic-" + unique,
            object_storage_access_id=environment["MINIO_ROOT_USER"],
            object_storage_access_secret=environment["MINIO_ROOT_PASSWORD"],
            clamd_host="127.0.0.1",
            clamd_port=scanner_port,
        )
        store = S3ObjectStore(settings)
        scanner = ClamAVScanner(settings)
        with httpx.Client(timeout=3, trust_env=False) as client:
            deadline = time.monotonic() + 60
            while True:
                try:
                    if (
                        client.get(
                            f"http://127.0.0.1:{storage_port}/minio/health/ready"
                        ).status_code
                        == 200
                    ):
                        break
                except httpx.TransportError:
                    pass
                if time.monotonic() > deadline:
                    raise RuntimeError("MinIO readiness failed")
                time.sleep(1)
        store.client.create_bucket(Bucket=settings.object_storage_bucket)
        print(f"Real MinIO ready; ClamAV image {clam_image}.", flush=True)
        try:
            wait_for_scanner(scanner)
        except RuntimeError:
            for command in (
                ["logs", "--tail", "120", clam],
                ["exec", clam, "tail", "-n", "80", "/var/log/clamav/clamd.log"],
            ):
                try:
                    print(docker(command), flush=True)
                except subprocess.CalledProcessError:
                    pass
            raise
        # Standard harmless antivirus test string exists only in memory and goes to localhost.
        eicar = b"X5O!P%@AP[4\\PZX54(P^)7CC)7}" + b"$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"
        try:
            scanner.scan(eicar)
        except IngestionError as exc:
            assert exc.code == "MALICIOUS_CONTENT"
        else:
            raise AssertionError("Real ClamAV did not reject its harmless standard test")
        print("Real ClamAV clean/malicious gate passed.", flush=True)
        engine = make_engine(settings)
        migrate(engine)
        with Session(engine) as db:
            db.add_all(
                [
                    Tenant(tenant_id="a", name="synthetic-a", active=True),
                    Tenant(tenant_id="b", name="synthetic-b", active=True),
                ]
            )
            db.flush()
            hashed = PASSWORDS.hash(password)
            for tenant in ("a", "b"):
                db.add(
                    User(
                        tenant_id=tenant,
                        user_id="reviewer",
                        username="reviewer",
                        password_hash=hashed,
                        role="reviewer",
                        active=True,
                    )
                )
            db.commit()
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        environment.update(
            APP_ENV="test",
            DATABASE_URL=settings.database_url.get_secret_value(),
            AUTH_SIGNING_KEY=key,
            OBJECT_STORAGE_MODE="s3",
            OBJECT_STORAGE_ENDPOINT=str(settings.object_storage_endpoint),
            OBJECT_STORAGE_BUCKET=settings.object_storage_bucket,
            OBJECT_STORAGE_ACCESS_ID=environment["MINIO_ROOT_USER"],
            OBJECT_STORAGE_ACCESS_SECRET=environment["MINIO_ROOT_PASSWORD"],
            PYTHONUTF8="1",
            PYTHONIOENCODING="utf-8",
        )
        server = subprocess.Popen(  # noqa: S603 - fixed local API; no shell
            [
                sys.executable,
                "-m",
                "bidding_compliance.api",
                "--host",
                "127.0.0.1",
                "--port",
                str(port),
            ],
            cwd=ROOT,
            env=environment,
            stdout=log,
            stderr=log,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        tokens = []
        with httpx.Client(
            base_url=f"http://127.0.0.1:{port}", timeout=20, trust_env=False
        ) as client:
            deadline = time.monotonic() + 20
            while True:
                try:
                    assert client.get("/api/v1/health").status_code == 200
                    break
                except httpx.TransportError:
                    if server.poll() is not None or time.monotonic() > deadline:
                        raise RuntimeError("API startup failed") from None
                    time.sleep(0.1)

            def login(tenant):
                reply = client.post(
                    "/api/v1/auth/token",
                    json=dict(tenant_id=tenant, username="reviewer", password=password),
                )
                assert reply.status_code == 200
                token = reply.json()["access_token"]
                tokens.append(token)
                return {"Authorization": "Bearer " + token}

            a, b = login("a"), login("b")
            project_response = client.post(
                "/api/v1/projects",
                headers={**a, "Idempotency-Key": str(uuid4())},
                json=dict(
                    name="synthetic",
                    procurement_type="government",
                    project_type="works",
                    region="synthetic",
                ),
            )
            assert project_response.status_code == 201
            project = project_response.json()["project_id"]
            for name, media, data in [
                ("synthetic.pdf", "application/pdf", pdf()),
                (
                    "synthetic.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    docx(),
                ),
                ("synthetic.png", "image/png", image()),
                ("synthetic-scan.pdf", "application/pdf", scanned_pdf()),
            ]:
                upload_path = f"/api/v1/projects/{project}/documents"
                headers = {**a, "Idempotency-Key": str(uuid4())}
                uploaded = client.post(
                    upload_path, headers=headers, files={"file": (name, data, media)}
                )
                assert uploaded.status_code == 202, uploaded.text
                document = uploaded.json()["document_id"]
                assert (
                    client.post(
                        upload_path, headers=headers, files={"file": (name, data, media)}
                    ).json()["document_id"]
                    == document
                )
                status = execute(engine, "a", document, parser, scanner, store, load_policy())
                assert status.status == "needs_confirmation", status.model_dump()
                path = upload_path + "/" + document
                fragments = client.get(path + "/sections", headers=a).json()["items"]
                assert fragments and all(f["source_hash"] == status.sha256 for f in fragments)
                if name.endswith(".png") or name == "synthetic-scan.pdf":
                    assert any(f["automatic"]["source"] == "ocr" for f in fragments)
                    assert (
                        "qualification"
                        in " ".join(f["effective"]["text"] for f in fragments).lower()
                    )
                    coordinate = "pdf_points" if name.endswith(".pdf") else "image_pixels"
                    assert all(
                        f["automatic"]["coordinate_system"] == coordinate
                        and f["automatic"]["page"] == 1
                        and f["automatic"]["bbox"] is not None
                        for f in fragments
                    )
                for suffix in ("", "/sections", "/context", "/original"):
                    assert client.get(path + suffix, headers=b).status_code == 404
                confirmation = client.post(
                    path + "/context",
                    headers=a,
                    json={
                        "expected_version": 1,
                        "reason": "synthetic reviewed confirmation",
                        "context": {
                            "name": "synthetic",
                            "procurement_type": "government",
                            "project_type": "works",
                            "region": "synthetic",
                        },
                    },
                )
                assert confirmation.status_code == 200, confirmation.text
                assert confirmation.json()["status"] == "ready"
                assert client.get(path + "/original", headers=a).content == data
                print("Real HTTP/S3/scan/isolated parse/context/download passed:", name, flush=True)
            unsafe = PdfWriter()
            unsafe.add_blank_page(width=100, height=100)
            unsafe.add_js("app.alert('synthetic');")
            buffer = io.BytesIO()
            unsafe.write(buffer)
            uploaded = client.post(
                f"/api/v1/projects/{project}/documents",
                headers={**a, "Idempotency-Key": str(uuid4())},
                files={"file": ("unsafe.pdf", buffer.getvalue(), "application/pdf")},
            )
            assert uploaded.status_code == 202
            result = execute(
                engine, "a", uploaded.json()["document_id"], parser, scanner, store, load_policy()
            )
            assert result.status == "rejected" and result.issues == ["MALICIOUS_CONTENT"]
        server.terminate()
        server.wait(timeout=10)
        log.flush()
        captured = (folder / "server.log").read_text(encoding="utf-8")
        assert all(
            value not in captured
            for value in [password, key, environment["MINIO_ROOT_PASSWORD"], *tokens]
        )
        print(
            "PASS: real file ingestion containers, S3, malware, OCR, "
            "locations, tenant isolation and HTTP",
            flush=True,
        )
    finally:
        if server is not None and server.poll() is None:
            server.terminate()
            server.wait(timeout=10)
        log.close()
        if engine is not None:
            engine.dispose()
        for name in reversed(names):
            docker(["rm", "--force", name])
        print("Stopped only the owned verification service/container resources.", flush=True)


if __name__ == "__main__":
    main()
