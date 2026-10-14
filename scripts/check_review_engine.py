"""PostgreSQL recovery/idempotency check using synthetic inputs and an offline gateway."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import selectors
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import psycopg
from docx import Document as WordDocument
from psycopg.rows import dict_row
from pydantic import ValidationError
from pypdf import PdfReader
from sqlalchemy import func, select, text
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session, sessionmaker

from bidding_compliance.core.config import Settings
from bidding_compliance.domain.identity import Principal, Role
from bidding_compliance.domain.rules.models import DiscoveryDocumentRequest
from bidding_compliance.domain.runtime_provenance import ReviewRuntimeProvenance
from bidding_compliance.ingestion.schemas import Fragment
from bidding_compliance.model_gateway.contracts import (
    DataPolicy,
    JsonObject,
    ModelPolicy,
    ModelRequest,
    ModelResponse,
    ModelUsage,
)
from bidding_compliance.model_gateway.errors import GatewayErrorCode, ModelGatewayError
from bidding_compliance.model_gateway.protocol import ModelGateway
from bidding_compliance.persistence.database import make_engine, set_tenant
from bidding_compliance.persistence.models import (
    AuditEvent,
    AuthSession,
    Document,
    DocumentContextVersion,
    DocumentSection,
    Project,
    ReviewArtifact,
    ReviewDecision,
    ReviewEngineRun,
    ReviewReportContent,
    ReviewResult,
    ReviewTask,
    User,
)
from bidding_compliance.services.identity import PASSWORDS
from bidding_compliance.services.reports import (
    citation_fragment,
    create_report,
    export_report,
    report_payload,
)
from bidding_compliance.services.review_engine import ReviewEngine, ReviewRequest
from bidding_compliance.services.review_persistence import (
    RESULT_PENDING_ID,
    approved_findings,
    authorized_run,
    publish_result,
    submit_decision,
)

ROOT = Path(__file__).resolve().parents[1]
REVIEWER_SESSION_ID = "task0009-review-verification-session"
PROFILE = "task0009-offline-verification"


def _run_async(coroutine):
    if sys.platform == "win32":
        return asyncio.run(
            coroutine,
            loop_factory=lambda: asyncio.SelectorEventLoop(selectors.SelectSelector()),
        )
    return asyncio.run(coroutine)


def _canonical_hash(value: object) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _policy() -> ModelPolicy:
    return ModelPolicy(
        profile=PROFILE,
        timeout_seconds=5.0,
        max_attempts=2,
        max_input_tokens=20_000,
        max_output_tokens=4_000,
        data_policy=DataPolicy.CONFIDENTIAL,
    )


class OfflineReviewGateway(ModelGateway):
    """Exercise the real gateway protocol without contacting a provider."""

    def __init__(self, *, fail_first: bool = False) -> None:
        self.fail_first = fail_first
        self.calls = 0
        self.saw_candidates = False
        self.saw_legal_results = False

    def prompt_sha256(self, prompt_version: str) -> str:
        return hashlib.sha256(("offline-test-prompt:" + prompt_version).encode()).hexdigest()

    async def generate_structured(
        self,
        request: ModelRequest,
        output_schema: JsonObject,
        policy: ModelPolicy,
    ) -> ModelResponse:
        self.calls += 1
        payload = dict(request.input_payload)
        candidates = payload["document_candidates"]["candidates"]
        retrieval = payload["legal_retrieval"]
        candidate_evidence = payload["candidate_legal_evidence"]
        if not candidates or not retrieval["results"]:
            raise AssertionError("gateway did not receive document and published-law results")
        if set(candidate_evidence) != {item["candidate_id"] for item in candidates}:
            raise AssertionError("gateway did not receive candidate-specific legal results")
        self.saw_candidates = True
        self.saw_legal_results = True
        if self.fail_first:
            self.fail_first = False
            raise ModelGatewayError(
                GatewayErrorCode.TIMEOUT,
                "synthetic transient provider timeout",
                retryable=True,
            )
        findings = []
        for candidate in candidates:
            evidence = candidate_evidence[candidate["candidate_id"]]
            legal_ids = [item["evidence_id"] for item in evidence.get("hits", [])]
            findings.append(
                {
                    "candidate_id": candidate["candidate_id"],
                    "assessment": "potential_issue",
                    "title": "Synthetic location requirement for human review",
                    "description": "Offline gateway output tied to deterministic evidence.",
                    "severity": "medium",
                    "source_quote_sha256": [
                        item["quote_sha256"] for item in candidate["source_excerpts"]
                    ],
                    "legal_evidence_ids": legal_ids,
                    "recommendation": "Review the cited source and applicable law.",
                    "reasoning_summary": "Synthetic protocol-conformance output.",
                    "confidence": 0.5,
                }
            )
        content = {"schema_version": "v1alpha1", "findings": findings}
        return ModelResponse(
            call_id=request.call_id,
            provider="offline-verification",
            model="synthetic-protocol-adapter",
            profile=policy.profile,
            content=content,
            usage=ModelUsage(input_tokens=1, output_tokens=1, estimated_cost=0.0),
            latency_ms=0,
            response_hash=_canonical_hash(content),
            finish_reason="stop",
        )


class NoCallGateway(ModelGateway):
    def prompt_sha256(self, prompt_version: str) -> str:
        return hashlib.sha256(("offline-test-prompt:" + prompt_version).encode()).hexdigest()

    async def generate_structured(
        self,
        request: ModelRequest,
        output_schema: JsonObject,
        policy: ModelPolicy,
    ) -> ModelResponse:
        raise AssertionError("recovered interrupt must not repeat model analysis")


def _principal(session_factory: sessionmaker[Session]) -> Principal:
    with session_factory() as session:
        auth = session.get(AuthSession, ("legal-a", REVIEWER_SESSION_ID))
        if auth is None or auth.revoked:
            raise AssertionError("synthetic recovery reviewer session is unavailable")
        return Principal("legal-a", "u", "u", Role.REVIEWER, REVIEWER_SESSION_ID)


def _resume_worker(review_task_id: str, run_id: str) -> None:
    database_url = os.environ.get("TASK0009_REVIEW_DATABASE_URL")
    if not database_url:
        raise RuntimeError("TASK0009_REVIEW_DATABASE_URL is required for the child worker")
    settings = Settings(
        _env_file=None,
        app_env="test",
        database_url=database_url,
        auth_signing_key=str(uuid4()) + str(uuid4()),
    )
    engine = make_engine(settings)
    factory = sessionmaker(engine, expire_on_commit=False)
    principal = _principal(factory)
    worker = ReviewEngine(
        settings=settings,
        session_factory=factory,
        gateway=NoCallGateway(),
        model_policy=_policy(),
        legal_root=ROOT,
    )
    result = _run_async(worker.resume(principal, review_task_id, run_id))
    if result.get("status") != "result_published":
        raise AssertionError("new worker did not publish the resumed review result")
    engine.dispose()


def verify_review_engine(
    admin_engine: Engine,
    runtime_engine: Engine,
    runtime_settings: Settings,
    principal: Principal,
    repository_root: Path,
) -> None:
    """Run retry, restart, human-resume and exactly-once publication against PostgreSQL."""
    if repository_root.resolve() != ROOT.resolve():
        raise ValueError("verification must run from this repository root")
    now = datetime.now(UTC)
    text_value = "投标人必须在本市注册。"
    source_bytes = b"synthetic TASK-0009 verification input"
    source_sha = hashlib.sha256(source_bytes).hexdigest()
    fragment = Fragment(
        locator="paragraph:0",
        coordinate_system="ooxml",
        kind="paragraph",
        section="资格要求",
        text=text_value,
        text_sha256=hashlib.sha256(text_value.encode("utf-8")).hexdigest(),
        source="native",
    )
    document_id = "task0009-doc-" + uuid4().hex[:12]
    with Session(admin_engine) as session:
        set_tenant(session, principal.tenant_id)
        session.execute(text("SELECT set_config('app.user_id', :actor, true)"), {"actor": "u"})
        session.add(
            Document(
                tenant_id=principal.tenant_id,
                document_id=document_id,
                project_id="p",
                status="ready",
                sha256=source_sha,
                object_ref="fixture://task0009/synthetic",
                file_format="docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                size_bytes=len(source_bytes),
                parse_version=1,
                context_version=1,
                issues=[],
            )
        )
        session.add(
            DocumentSection(
                tenant_id=principal.tenant_id,
                document_id=document_id,
                section_id="section-1",
                ordinal=0,
                automatic=fragment.model_dump(mode="json"),
                revision=1,
            )
        )
        session.add(
            DocumentContextVersion(
                tenant_id=principal.tenant_id,
                document_id=document_id,
                version=1,
                payload={
                    "document_id": document_id,
                    "version": 1,
                    "candidates": [],
                    "confirmed": True,
                },
                changed_by="u",
                changed_at=now,
            )
        )
        session.add(
            AuthSession(
                tenant_id=principal.tenant_id,
                session_id=REVIEWER_SESSION_ID,
                user_id="u",
                expires_at=now + timedelta(hours=1),
                revoked=False,
            )
        )
        session.commit()

    policy = _policy()
    gateway = OfflineReviewGateway(fail_first=True)
    worker = ReviewEngine(
        settings=runtime_settings,
        session_factory=sessionmaker(runtime_engine, expire_on_commit=False),
        gateway=gateway,
        model_policy=policy,
        legal_root=repository_root,
    )
    run_id = str(uuid4())
    request = ReviewRequest(
        review_task_id="task0009-" + uuid4().hex[:20],
        project_id="p",
        run_id=run_id,
        document_requests=(
            DiscoveryDocumentRequest(
                document_id=document_id,
                raw_source_sha256=source_sha,
                parse_version=1,
                context_version=1,
            ),
        ),
        review_date=date.today(),
        prompt_version="task0009-offline-v1",
        model_profile=PROFILE,
    )
    failed = _run_async(worker.start(principal, request))
    if failed.get("status") != "failed" or not failed.get("errors"):
        raise AssertionError("transient model failure did not enter the durable failed state")
    failure_errors = failed["errors"]
    if not any(
        error.get("code") == GatewayErrorCode.TIMEOUT.value
        and error.get("node") == "semantic_analysis"
        for error in failure_errors
    ):
        raise AssertionError(
            "offline timeout injection did not reach semantic_analysis: "
            f"gateway_calls={gateway.calls}, errors={failure_errors!r}"
        )
    with Session(runtime_engine) as session:
        failed_run = authorized_run(session, principal, request.review_task_id, request.run_id)
        if (
            failed_run.status != "failed"
            or failed_run.failure_code != GatewayErrorCode.TIMEOUT.value
            or failed_run.retry_count != 0
        ):
            raise AssertionError(
                "durable timeout classification is inconsistent: "
                f"status={failed_run.status!r}, "
                f"failure_code={failed_run.failure_code!r}, "
                f"retry_count={failed_run.retry_count!r}"
            )
        try:
            provenance = ReviewRuntimeProvenance.model_validate_json(
                json.dumps(failed_run.snapshot["runtime_provenance"], ensure_ascii=False)
            )
        except (KeyError, ValidationError) as exc:
            raise AssertionError("review run did not persist valid runtime provenance") from exc
        if (
            len(provenance.parser_documents) != 1
            or provenance.parser_documents[0].document_id != document_id
            or provenance.parser_documents[0].identity is not None
            or provenance.embedding_identity.provider != "character"
            or provenance.vector_index_build_id is not None
            or provenance.prompt_version != request.prompt_version
            or provenance.model_profile != request.model_profile
            or provenance.model_config_revision_id
            != (failed_run.snapshot.get("configuration_revision_id") or "code-default")
            or provenance.model_config_sha256
            != failed_run.snapshot["configuration_snapshot_hash"]
            or failed_run.snapshot_hash != _canonical_hash(failed_run.snapshot)
        ):
            raise AssertionError("persisted review runtime provenance is not bound to its run")
    pending = _run_async(worker.retry(principal, request.review_task_id, request.run_id))
    if pending.get("status") != "waiting_human":
        raise AssertionError("bounded retry did not reach the persisted human interrupt")
    if gateway.calls != 2 or not gateway.saw_candidates or not gateway.saw_legal_results:
        raise AssertionError("gateway input or retry boundary was not exercised")
    with Session(runtime_engine) as session:
        authorized_run(session, principal, request.review_task_id, request.run_id)
        merged = session.scalar(
            select(ReviewArtifact).where(
                ReviewArtifact.tenant_id == principal.tenant_id,
                ReviewArtifact.review_task_id == request.review_task_id,
                ReviewArtifact.run_id == request.run_id,
                ReviewArtifact.kind == "merged_findings",
            )
        )
        if merged is None or not merged.content["findings"]:
            raise AssertionError("synthetic evidence review produced no merged findings")
        source_finding = merged.content["findings"][0]
        for finding in merged.content["findings"]:
            law_evidence = [
                item for item in finding["evidence"] if item["kind"] in {"law", "policy"}
            ]
            if finding["evidence_status"] == "confirmed":
                raise AssertionError("model output bypassed independent evidence downgrade")
            if not law_evidence and finding["evidence_status"] != "insufficient_evidence":
                raise AssertionError("missing legal evidence did not produce insufficient_evidence")

    restarted_worker = ReviewEngine(
        settings=runtime_settings,
        session_factory=sessionmaker(runtime_engine, expire_on_commit=False),
        gateway=NoCallGateway(),
        model_policy=policy,
        legal_root=repository_root,
    )
    duplicate_delivery = _run_async(restarted_worker.start(principal, request))
    if duplicate_delivery.get("status") != "waiting_human":
        raise AssertionError("duplicate delivery did not return the persisted interrupt")
    pending_findings = pending.get("findings", [])
    with Session(runtime_engine) as session:
        waiting_row = authorized_run(session, principal, request.review_task_id, request.run_id)
        merged = session.scalar(
            select(ReviewArtifact).where(
                ReviewArtifact.tenant_id == principal.tenant_id,
                ReviewArtifact.review_task_id == request.review_task_id,
                ReviewArtifact.run_id == request.run_id,
                ReviewArtifact.kind == "merged_findings",
            )
        )
        if merged is None:
            raise AssertionError("merged evidence artifact is unavailable before review")
        try:
            publish_result(
                session,
                principal,
                row=waiting_row,
                findings=merged.content["findings"],
            )
        except PermissionError:
            session.rollback()
        else:
            raise AssertionError("result was published without a persisted human decision")
    if pending_findings:
        finding_id = pending_findings[0]["finding_id"]
        decision = "accept"
    else:
        finding_id = RESULT_PENDING_ID
        decision = "confirm_no_findings"
    restarted_worker.submit_decision(
        principal,
        review_task_id=request.review_task_id,
        run_id=request.run_id,
        finding_id=finding_id,
        decision=decision,
        request_id="decision-" + uuid4().hex[:20],
        reason="Synthetic PostgreSQL recovery acceptance test.",
    )

    child_environment = os.environ.copy()
    child_environment["TASK0009_REVIEW_DATABASE_URL"] = (
        runtime_settings.database_url.get_secret_value()
        if runtime_settings.database_url is not None
        else ""
    )
    script = Path(__file__).resolve()
    resumed = subprocess.run(  # noqa: S603 - fixed local script and UUID/validated task arguments.
        [sys.executable, str(script), "--resume", request.review_task_id, request.run_id],
        cwd=repository_root,
        env=child_environment,
        capture_output=True,
        text=True,
        check=False,
    )
    if resumed.returncode != 0:
        raise AssertionError(
            "new-process recovery failed: " + (resumed.stderr or resumed.stdout)[-4000:]
        )

    factory = sessionmaker(runtime_engine, expire_on_commit=False)
    with factory() as session:
        row = authorized_run(session, principal, request.review_task_id, request.run_id)
        if row.status != "result_published":
            raise AssertionError("durable run did not reach result_published")
        results = list(
            session.scalars(
                select(ReviewResult).where(
                    ReviewResult.tenant_id == principal.tenant_id,
                    ReviewResult.review_task_id == request.review_task_id,
                    ReviewResult.run_id == request.run_id,
                )
            )
        )
        if len(results) != 1:
            raise AssertionError("review result publication was not exactly once")
        persisted_decisions = list(
            session.scalars(
                select(ReviewDecision).where(
                    ReviewDecision.tenant_id == principal.tenant_id,
                    ReviewDecision.review_task_id == request.review_task_id,
                    ReviewDecision.run_id == request.run_id,
                )
            )
        )
        if len(persisted_decisions) != 1:
            raise AssertionError("human decision history was not persisted exactly once")
        decision_record = persisted_decisions[0]
        if (
            decision_record.before_state is None
            or decision_record.after_state is None
            or decision_record.reason.strip() == ""
            or decision_record.request_id == ""
        ):
            raise AssertionError("human review did not retain before/after values and reason")
        if decision_record.finding_id == RESULT_PENDING_ID:
            if decision_record.after_state.get("review_status") != "confirmed_no_findings":
                raise AssertionError("empty result confirmation history is incomplete")
        elif (
            decision_record.before_state.get("review_status") != "pending"
            or decision_record.after_state.get("review_status") != "accepted"
            or decision_record.before_state.get("evidence_status")
            != decision_record.after_state.get("evidence_status")
        ):
            raise AssertionError("human status changed or omitted the independent evidence state")
        published = results[0]
        first_docx = create_report(
            session,
            principal,
            result_id=published.result_id,
            output_format="docx",
            idempotency_key="task0010-postgres-docx-first",
        )
        repeated_docx = create_report(
            session,
            principal,
            result_id=published.result_id,
            output_format="docx",
            idempotency_key="task0010-postgres-docx-first",
        )
        pdf_report = create_report(
            session,
            principal,
            result_id=published.result_id,
            output_format="pdf",
            idempotency_key="task0010-postgres-pdf-first",
        )
        if first_docx.report_id != repeated_docx.report_id:
            raise AssertionError("report generation retry created a second formal artifact")
        if (
            first_docx.report_content_id != pdf_report.report_content_id
            or first_docx.report_content_id != repeated_docx.report_content_id
        ):
            raise AssertionError("Word and PDF did not reuse one structured report version")
        set_tenant(session, principal.tenant_id)
        structured_report = session.get(
            ReviewReportContent,
            (principal.tenant_id, first_docx.report_content_id),
        )
        if (
            structured_report is None
            or structured_report.content["report_id"] != first_docx.report_content_id
            or structured_report.content_hash == ""
        ):
            raise AssertionError("structured report version was not durably registered")
        if (
            first_docx.snapshot_hash != published.snapshot_hash
            or pdf_report.snapshot_hash != published.snapshot_hash
        ):
            raise AssertionError("Word/PDF did not bind to the immutable result snapshot")
        if not first_docx.content.startswith(b"PK") or not pdf_report.content.startswith(b"%PDF-"):
            raise AssertionError("report artifacts do not contain their declared file formats")
        with ZipFile(BytesIO(first_docx.content)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise AssertionError("Word artifact is not a parseable DOCX package")
        if not WordDocument(BytesIO(first_docx.content)).paragraphs:
            raise AssertionError("Word report has no readable body")
        if len(PdfReader(BytesIO(pdf_report.content)).pages) < 1:
            raise AssertionError("PDF report has no pages")
        if published.content["findings"]:
            report_body = WordDocument(BytesIO(first_docx.content))
            if not any("投标人必须在本市注册" in p.text for p in report_body.paragraphs):
                raise AssertionError("report omitted its source excerpt")
            payload = report_payload(session, principal, published)
            cite = payload["findings"][0]["citations"][0]
            resolved = citation_fragment(
                session,
                principal,
                project_id=published.content["project_id"],
                document_id=cite["document_id"],
                section_id=cite["section_id"],
                revision=cite["section_revision"],
                source_hash=cite["raw_source_sha256"],
            )
            if resolved["fragment"]["text"] != cite["quote"]:
                raise AssertionError(
                    "report citation did not resolve to the exact original fragment"
                )
        exported = export_report(session, principal, first_docx.report_id)
        if exported.content_hash != first_docx.content_hash:
            raise AssertionError("authorized report export changed the artifact")
        set_tenant(session, principal.tenant_id)
        exported_events = session.scalar(
            select(func.count(AuditEvent.event_id)).where(
                AuditEvent.tenant_id == principal.tenant_id,
                AuditEvent.actor_id == principal.user_id,
                AuditEvent.action == "report.exported",
                AuditEvent.resource_id == first_docx.report_id,
            )
        )
        if exported_events != 1:
            raise AssertionError("report export did not produce exactly one audit event")
        if pending_findings:
            concurrent_reviewer = Principal(
                principal.tenant_id, "u2", "reviewer-two", Role.REVIEWER, str(uuid4())
            )
            auditor = Principal(principal.tenant_id, "u3", "auditor", Role.AUDITOR, str(uuid4()))
            tenant_admin = Principal(
                principal.tenant_id, "u4", "tenant-admin", Role.ADMIN, str(uuid4())
            )
            with Session(admin_engine) as admin_session:
                for secondary, role in (
                    (concurrent_reviewer, "reviewer"),
                    (auditor, "auditor"),
                    (tenant_admin, "admin"),
                ):
                    admin_session.add(
                        User(
                            tenant_id=principal.tenant_id,
                            user_id=secondary.user_id,
                            username="synthetic-" + secondary.user_id,
                            password_hash=PASSWORDS.hash(str(uuid4())),
                            role=role,
                            active=True,
                        )
                    )
                admin_session.flush()
                for secondary, _ in (
                    (concurrent_reviewer, "reviewer"),
                    (auditor, "auditor"),
                    (tenant_admin, "admin"),
                ):
                    admin_session.add(
                        AuthSession(
                            tenant_id=principal.tenant_id,
                            session_id=secondary.session_id,
                            user_id=secondary.user_id,
                            expires_at=datetime.now(UTC) + timedelta(hours=1),
                            revoked=False,
                        )
                    )
                admin_session.commit()

            def create_decision_case(label: str) -> tuple[str, str, str]:
                synthetic_task = f"task0010-{label}-" + uuid4().hex[:12]
                synthetic_run = str(uuid4())
                snapshot = json.loads(json.dumps(row.snapshot))
                snapshot.update(
                    {
                        "tenant_id": principal.tenant_id,
                        "project_id": row.project_id,
                        "review_task_id": synthetic_task,
                        "run_id": synthetic_run,
                    }
                )
                snapshot_digest = _canonical_hash(snapshot)
                artifact_content = {"findings": [source_finding]}
                artifact_digest = _canonical_hash(artifact_content)
                artifact_id = hashlib.sha256(
                    (
                        f"{principal.tenant_id}\0{synthetic_task}\0{synthetic_run}\0"
                        f"merged_findings\0{artifact_digest}"
                    ).encode()
                ).hexdigest()
                with Session(admin_engine) as admin_session:
                    admin_session.add(
                        ReviewTask(
                            tenant_id=principal.tenant_id,
                            review_task_id=synthetic_task,
                            project_id=row.project_id,
                            run_id=synthetic_run,
                            status="waiting_human",
                        )
                    )
                    admin_session.flush()
                    admin_session.add(
                        ReviewEngineRun(
                            tenant_id=principal.tenant_id,
                            review_task_id=synthetic_task,
                            run_id=synthetic_run,
                            project_id=row.project_id,
                            actor_id=principal.user_id,
                            status="waiting_human",
                            snapshot=snapshot,
                            snapshot_hash=snapshot_digest,
                            request_hash="d" * 64,
                            retry_count=0,
                            publication_id=row.publication_id,
                            revision_id=row.revision_id,
                            created_at=datetime.now(UTC),
                            updated_at=datetime.now(UTC),
                        )
                    )
                    admin_session.flush()
                    admin_session.add(
                        ReviewArtifact(
                            tenant_id=principal.tenant_id,
                            artifact_id=artifact_id,
                            review_task_id=synthetic_task,
                            run_id=synthetic_run,
                            kind="merged_findings",
                            content=artifact_content,
                            content_hash=artifact_digest,
                            created_at=datetime.now(UTC),
                        )
                    )
                    admin_session.commit()
                return synthetic_task, synthetic_run, snapshot_digest

            for action, reviewer in (
                ("accept", principal),
                ("modify", concurrent_reviewer),
                ("reject", principal),
            ):
                case_task, case_run, case_hash = create_decision_case(action)
                key = "t10-" + uuid4().hex
                changes = {"title": "经人工修改的风险标题"} if action == "modify" else None
                with factory() as decision_session:
                    stored = submit_decision(
                        decision_session,
                        reviewer,
                        review_task_id=case_task,
                        run_id=case_run,
                        finding_id=source_finding["finding_id"],
                        decision=action,
                        request_id=key,
                        reason="Synthetic TASK-0010 action verification.",
                        modified_finding=changes,
                        expected_snapshot_hash=case_hash,
                        request_trace_id=str(uuid4()),
                    )
                    duplicate = submit_decision(
                        decision_session,
                        reviewer,
                        review_task_id=case_task,
                        run_id=case_run,
                        finding_id=source_finding["finding_id"],
                        decision=action,
                        request_id=key,
                        reason="Synthetic TASK-0010 action verification.",
                        modified_finding=changes,
                        expected_snapshot_hash=case_hash,
                    )
                    projected = approved_findings(
                        decision_session,
                        reviewer,
                        authorized_run(decision_session, reviewer, case_task, case_run),
                        [source_finding],
                    )[0]
                    if (
                        stored.decision_id != duplicate.decision_id
                        or stored.before_state.get("review_status") != "pending"
                        or stored.after_state.get("review_status") != projected["review_status"]
                        or stored.before_state.get("evidence_status")
                        != stored.after_state.get("evidence_status")
                        or stored.reason.strip() == ""
                        or stored.request_trace_id is None
                    ):
                        raise AssertionError(
                            f"{action} review did not persist an immutable audit pair"
                        )
                    if action == "modify" and projected["title"] != changes["title"]:
                        raise AssertionError(
                            "human modification was not included in the result projection"
                        )

            race_task, race_run, race_hash = create_decision_case("race")

            def race_decision(pair):
                reviewer, decision_value = pair
                try:
                    with factory() as decision_session:
                        return submit_decision(
                            decision_session,
                            reviewer,
                            review_task_id=race_task,
                            run_id=race_run,
                            finding_id=source_finding["finding_id"],
                            decision=decision_value,
                            request_id="t10-" + uuid4().hex,
                            reason="Concurrent TASK-0010 reviewer conflict check.",
                            expected_snapshot_hash=race_hash,
                        )
                except ValueError:
                    return None

            with ThreadPoolExecutor(max_workers=2) as pool:
                competing = list(
                    pool.map(
                        race_decision, ((principal, "accept"), (concurrent_reviewer, "reject"))
                    )
                )
            if sum(item is not None for item in competing) != 1:
                raise AssertionError("concurrent reviewers both changed the same risk")

            denied_task, denied_run, denied_hash = create_decision_case("denied")
            with factory() as decision_session:
                try:
                    submit_decision(
                        decision_session,
                        auditor,
                        review_task_id=denied_task,
                        run_id=denied_run,
                        finding_id=source_finding["finding_id"],
                        decision="accept",
                        request_id="t10-" + uuid4().hex,
                        reason="Non-reviewer denial check.",
                        expected_snapshot_hash=denied_hash,
                    )
                except PermissionError:
                    pass
                else:
                    raise AssertionError("non-reviewer submitted a human review decision")
                set_tenant(decision_session, principal.tenant_id)
                denied_event = decision_session.scalar(
                    select(func.count(AuditEvent.event_id)).where(
                        AuditEvent.tenant_id == principal.tenant_id,
                        AuditEvent.action == "review.decision_denied",
                    )
                )
                if not denied_event:
                    raise AssertionError("denied review decision was not audited")

            with factory() as export_session:
                try:
                    export_report(export_session, tenant_admin, first_docx.report_id)
                except PermissionError:
                    pass
                else:
                    raise AssertionError("tenant admin exported a reviewer-only report")
                set_tenant(export_session, principal.tenant_id)
                export_denied = export_session.scalar(
                    select(func.count(AuditEvent.event_id)).where(
                        AuditEvent.tenant_id == principal.tenant_id,
                        AuditEvent.actor_id == tenant_admin.user_id,
                        AuditEvent.action == "report.export_denied",
                        AuditEvent.resource_id == first_docx.report_id,
                    )
                )
                if export_denied != 1:
                    raise AssertionError("denied report export was not audited")

    def duplicate_publish(_):
        with factory() as session:
            row = authorized_run(session, principal, request.review_task_id, request.run_id)
            merged = session.scalar(
                select(ReviewArtifact).where(
                    ReviewArtifact.tenant_id == principal.tenant_id,
                    ReviewArtifact.review_task_id == request.review_task_id,
                    ReviewArtifact.run_id == request.run_id,
                    ReviewArtifact.kind == "merged_findings",
                )
            )
            if merged is None:
                raise AssertionError("merged evidence artifact is unavailable")
            return publish_result(
                session,
                principal,
                row=row,
                findings=merged.content["findings"],
            ).result_id

    with ThreadPoolExecutor(max_workers=4) as pool:
        duplicate_ids = list(pool.map(duplicate_publish, range(4)))
    if any(result_id != results[0].result_id for result_id in duplicate_ids):
        raise AssertionError("concurrent duplicate publication created a second result")
    with factory() as session:
        authorized_run(session, principal, request.review_task_id, request.run_id)
        result_count = session.scalar(
            select(func.count(ReviewResult.result_id)).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == request.review_task_id,
                ReviewResult.run_id == request.run_id,
            )
        )
        if result_count != 1:
            raise AssertionError("concurrent publication persisted multiple result versions")

    delivered_again = _run_async(restarted_worker.start(principal, request))
    if delivered_again.get("status") != "result_published" or gateway.calls != 2:
        raise AssertionError("post-completion duplicate delivery reran analysis")

    concurrent_request = ReviewRequest(
        review_task_id="task0009-concurrent-" + uuid4().hex[:12],
        project_id=request.project_id,
        run_id=str(uuid4()),
        document_requests=request.document_requests,
        review_date=request.review_date,
        prompt_version=request.prompt_version,
        model_profile=request.model_profile,
    )
    concurrent_gateway = OfflineReviewGateway()
    concurrent_worker = ReviewEngine(
        settings=runtime_settings,
        session_factory=factory,
        gateway=concurrent_gateway,
        model_policy=policy,
        legal_root=repository_root,
    )
    concurrent_pending = _run_async(concurrent_worker.start(principal, concurrent_request))
    if concurrent_pending.get("status") != "waiting_human":
        raise AssertionError("concurrent recovery scenario missed the persisted interrupt")
    concurrent_findings = concurrent_pending.get("findings", [])
    concurrent_finding_id = (
        concurrent_findings[0]["finding_id"] if concurrent_findings else RESULT_PENDING_ID
    )
    concurrent_decision = "accept" if concurrent_findings else "confirm_no_findings"
    concurrent_worker.submit_decision(
        principal,
        review_task_id=concurrent_request.review_task_id,
        run_id=concurrent_request.run_id,
        finding_id=concurrent_finding_id,
        decision=concurrent_decision,
        request_id="concurrent-decision-" + uuid4().hex[:12],
        reason="Synthetic concurrent recovery verification.",
    )

    worker_a = ReviewEngine(
        settings=runtime_settings,
        session_factory=factory,
        gateway=NoCallGateway(),
        model_policy=policy,
        legal_root=repository_root,
    )
    worker_b = ReviewEngine(
        settings=runtime_settings,
        session_factory=factory,
        gateway=NoCallGateway(),
        model_policy=policy,
        legal_root=repository_root,
    )
    with factory() as session:
        authorized_run(
            session, principal, concurrent_request.review_task_id, concurrent_request.run_id
        )
        artifact_count_before = session.scalar(
            select(func.count(ReviewArtifact.artifact_id)).where(
                ReviewArtifact.tenant_id == principal.tenant_id,
                ReviewArtifact.review_task_id == concurrent_request.review_task_id,
                ReviewArtifact.run_id == concurrent_request.run_id,
            )
        )
        decision_count_before = session.scalar(
            select(func.count(ReviewDecision.decision_id)).where(
                ReviewDecision.tenant_id == principal.tenant_id,
                ReviewDecision.review_task_id == concurrent_request.review_task_id,
                ReviewDecision.run_id == concurrent_request.run_id,
            )
        )

    async def race_recovery_calls():
        barrier = asyncio.Barrier(2)

        async def resume_once(worker):
            await barrier.wait()
            try:
                return await worker.resume(
                    principal,
                    concurrent_request.review_task_id,
                    concurrent_request.run_id,
                )
            except ValueError as exc:
                if str(exc) in {
                    "review is not waiting for human review",
                    "durable human interrupt is not available to resume",
                }:
                    return {"already_processed": True}
                raise

        return await asyncio.gather(resume_once(worker_a), resume_once(worker_b))

    concurrent_outcomes = _run_async(race_recovery_calls())
    published_outcomes = [
        item for item in concurrent_outcomes if item.get("status") == "result_published"
    ]
    safe_duplicate_outcomes = [
        item
        for item in concurrent_outcomes
        if item.get("busy") is True or item.get("already_processed") is True
    ]
    if len(published_outcomes) != 1 or len(safe_duplicate_outcomes) != 1:
        raise AssertionError(
            "concurrent recovery did not yield one publisher and one safe duplicate response"
        )
    with factory() as session:
        concurrent_run = authorized_run(
            session, principal, concurrent_request.review_task_id, concurrent_request.run_id
        )
        result_count = session.scalar(
            select(func.count(ReviewResult.result_id)).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == concurrent_request.review_task_id,
                ReviewResult.run_id == concurrent_request.run_id,
            )
        )
        artifact_count_after = session.scalar(
            select(func.count(ReviewArtifact.artifact_id)).where(
                ReviewArtifact.tenant_id == principal.tenant_id,
                ReviewArtifact.review_task_id == concurrent_request.review_task_id,
                ReviewArtifact.run_id == concurrent_request.run_id,
            )
        )
        decision_count_after = session.scalar(
            select(func.count(ReviewDecision.decision_id)).where(
                ReviewDecision.tenant_id == principal.tenant_id,
                ReviewDecision.review_task_id == concurrent_request.review_task_id,
                ReviewDecision.run_id == concurrent_request.run_id,
            )
        )
        if (
            result_count != 1
            or concurrent_run.status != "result_published"
            or artifact_count_after != artifact_count_before
            or decision_count_after != decision_count_before
        ):
            raise AssertionError("concurrent recovery persisted multiple result versions")
    checkpoint_state = _run_async(worker_a.start(principal, concurrent_request))
    if checkpoint_state.get("status") != "result_published":
        raise AssertionError("reopened worker did not observe the completed checkpoint state")

    failure_request = ReviewRequest(
        review_task_id="task0009-fault-" + uuid4().hex[:16],
        project_id=request.project_id,
        run_id=str(uuid4()),
        document_requests=request.document_requests,
        review_date=request.review_date,
        prompt_version=request.prompt_version,
        model_profile=request.model_profile,
    )
    failure_gateway = OfflineReviewGateway()
    failure_worker = ReviewEngine(
        settings=runtime_settings,
        session_factory=factory,
        gateway=failure_gateway,
        model_policy=policy,
        legal_root=repository_root,
    )
    failure_pending = _run_async(failure_worker.start(principal, failure_request))
    if failure_pending.get("status") != "waiting_human":
        raise AssertionError("publication failure scenario missed human review")
    failure_findings = failure_pending.get("findings", [])
    failure_finding_id = (
        failure_findings[0]["finding_id"] if failure_findings else RESULT_PENDING_ID
    )
    failure_decision = "accept" if failure_findings else "confirm_no_findings"
    failure_worker.submit_decision(
        principal,
        review_task_id=failure_request.review_task_id,
        run_id=failure_request.run_id,
        finding_id=failure_finding_id,
        decision=failure_decision,
        request_id="fault-decision-" + uuid4().hex[:12],
        reason="Synthetic publication failure verification.",
    )
    with factory() as session:
        failure_row = authorized_run(
            session, principal, failure_request.review_task_id, failure_request.run_id
        )
        failure_artifact = session.scalar(
            select(ReviewArtifact).where(
                ReviewArtifact.tenant_id == principal.tenant_id,
                ReviewArtifact.review_task_id == failure_request.review_task_id,
                ReviewArtifact.run_id == failure_request.run_id,
                ReviewArtifact.kind == "merged_findings",
            )
        )
        if failure_artifact is None:
            raise AssertionError("fault-injection evidence artifact is unavailable")
        failure_findings_payload = failure_artifact.content["findings"]

    with admin_engine.begin() as connection:
        connection.execute(
            text("""
                CREATE FUNCTION task0009_fail_review_result_insert()
                RETURNS trigger LANGUAGE plpgsql AS $$
                BEGIN
                    RAISE EXCEPTION 'synthetic TASK-0009 publication database failure';
                END;
                $$
            """)
        )
        connection.execute(
            text("""
                CREATE TRIGGER task0009_fail_review_result_insert
                BEFORE INSERT ON review_result
                FOR EACH ROW EXECUTE FUNCTION task0009_fail_review_result_insert()
            """)
        )
    try:
        with factory() as session:
            failure_row = authorized_run(
                session, principal, failure_request.review_task_id, failure_request.run_id
            )
            try:
                publish_result(
                    session,
                    principal,
                    row=failure_row,
                    findings=failure_findings_payload,
                )
            except Exception as exc:
                if "synthetic TASK-0009 publication database failure" not in str(exc):
                    raise
                session.rollback()
            else:
                raise AssertionError("database publication failure injection did not fire")
    finally:
        with admin_engine.begin() as connection:
            connection.execute(
                text("DROP TRIGGER task0009_fail_review_result_insert ON review_result")
            )
            connection.execute(text("DROP FUNCTION task0009_fail_review_result_insert()"))
    with factory() as session:
        authorized_run(session, principal, failure_request.review_task_id, failure_request.run_id)
        result_count = session.scalar(
            select(func.count(ReviewResult.result_id)).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == failure_request.review_task_id,
                ReviewResult.run_id == failure_request.run_id,
            )
        )
        failed_publication = authorized_run(
            session, principal, failure_request.review_task_id, failure_request.run_id
        )
        if result_count != 0 or failed_publication.status != "waiting_human":
            raise AssertionError("failed publication left a partial or advanced result")

    with factory() as session:
        failure_row = authorized_run(
            session, principal, failure_request.review_task_id, failure_request.run_id
        )
        original_commit = session.commit

        def commit_then_lose_acknowledgement():
            original_commit()
            raise ConnectionError("synthetic commit acknowledgement lost")

        session.commit = commit_then_lose_acknowledgement
        try:
            publish_result(
                session,
                principal,
                row=failure_row,
                findings=failure_findings_payload,
            )
        except ConnectionError as exc:
            if str(exc) != "synthetic commit acknowledgement lost":
                raise
        else:
            raise AssertionError("commit acknowledgement loss injection did not fire")

    with factory() as session:
        committed_run = authorized_run(
            session, principal, failure_request.review_task_id, failure_request.run_id
        )
        committed_result = session.scalar(
            select(ReviewResult).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == failure_request.review_task_id,
                ReviewResult.run_id == failure_request.run_id,
            )
        )
        if committed_run.status != "result_published" or committed_result is None:
            raise AssertionError("transaction did not persist before losing its acknowledgement")
        committed_result_id = committed_result.result_id

    with factory() as session:
        failure_row = authorized_run(
            session, principal, failure_request.review_task_id, failure_request.run_id
        )
        recovered_result = publish_result(
            session,
            principal,
            row=failure_row,
            findings=failure_findings_payload,
        )
        recovered_result_id = recovered_result.result_id
    with factory() as session:
        authorized_run(session, principal, failure_request.review_task_id, failure_request.run_id)
        result_count = session.scalar(
            select(func.count(ReviewResult.result_id)).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == failure_request.review_task_id,
                ReviewResult.run_id == failure_request.run_id,
            )
        )
        persisted_result = session.scalar(
            select(ReviewResult).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == failure_request.review_task_id,
                ReviewResult.run_id == failure_request.run_id,
            )
        )
        if (
            result_count != 1
            or persisted_result is None
            or persisted_result.result_id != committed_result_id
            or persisted_result.result_id != recovered_result_id
        ):
            raise AssertionError("lost-acknowledgement retry duplicated the result")

    cancelled_request = ReviewRequest(
        review_task_id="task0009-cancel-" + uuid4().hex[:16],
        project_id=request.project_id,
        run_id=str(uuid4()),
        document_requests=request.document_requests,
        review_date=request.review_date,
        prompt_version=request.prompt_version,
        model_profile=request.model_profile,
    )
    worker.create_run(principal, cancelled_request)
    worker.cancel(principal, cancelled_request.review_task_id, cancelled_request.run_id)
    with factory() as session:
        cancelled = authorized_run(
            session, principal, cancelled_request.review_task_id, cancelled_request.run_id
        )
        if cancelled.status != "cancelled":
            raise AssertionError("queued review cancellation was not persisted")

    changed_request = ReviewRequest(
        review_task_id="task0009-snapshot-" + uuid4().hex[:16],
        project_id=request.project_id,
        run_id=str(uuid4()),
        document_requests=request.document_requests,
        review_date=request.review_date,
        prompt_version=request.prompt_version,
        model_profile=request.model_profile,
    )
    snapshot_worker = ReviewEngine(
        settings=runtime_settings,
        session_factory=sessionmaker(runtime_engine, expire_on_commit=False),
        gateway=OfflineReviewGateway(),
        model_policy=policy,
        legal_root=repository_root,
    )
    changed_pending = _run_async(snapshot_worker.start(principal, changed_request))
    if changed_pending.get("status") != "waiting_human":
        raise AssertionError("snapshot recovery scenario did not reach the human interrupt")
    changed_findings = changed_pending.get("findings", [])
    changed_finding_id = (
        changed_findings[0]["finding_id"] if changed_findings else RESULT_PENDING_ID
    )
    snapshot_worker.submit_decision(
        principal,
        review_task_id=changed_request.review_task_id,
        run_id=changed_request.run_id,
        finding_id=changed_finding_id,
        decision="accept" if changed_findings else "confirm_no_findings",
        request_id="snapshot-decision-" + uuid4().hex[:16],
        reason="Synthetic snapshot invalidation verification.",
    )

    with Session(admin_engine) as session:
        auth = session.get(AuthSession, (principal.tenant_id, principal.session_id))
        if auth is None:
            raise AssertionError("reviewer session disappeared before revocation check")
        auth.revoked = True
        session.commit()
    try:
        _run_async(
            snapshot_worker.resume(
                principal, changed_request.review_task_id, changed_request.run_id
            )
        )
    except PermissionError:
        pass
    else:
        raise AssertionError("revoked reviewer identity resumed a review")
    finally:
        with Session(admin_engine) as session:
            auth = session.get(AuthSession, (principal.tenant_id, principal.session_id))
            if auth is not None:
                auth.revoked = False
                session.commit()

    with Session(admin_engine) as session:
        project = session.get(Project, (principal.tenant_id, request.project_id))
        if project is None:
            raise AssertionError("synthetic project disappeared before snapshot check")
        original_version = project.version
        original_context = dict(project.context)
        project.version += 1
        project.context = {**original_context, "name": "synthetic-context-changed"}
        session.commit()
    try:
        _run_async(
            snapshot_worker.resume(
                principal, changed_request.review_task_id, changed_request.run_id
            )
        )
    except ValueError as exc:
        if "input changed" not in str(exc) and "context changed" not in str(exc):
            raise
    else:
        raise AssertionError("changed project snapshot resumed under its old run")
    finally:
        with Session(admin_engine) as session:
            project = session.get(Project, (principal.tenant_id, request.project_id))
            if project is not None:
                project.version = original_version
                project.context = original_context
                session.commit()
    with factory() as session:
        authorized_run(
            session,
            principal,
            changed_request.review_task_id,
            changed_request.run_id,
        )
        stale_result_count = session.scalar(
            select(func.count(ReviewResult.result_id)).where(
                ReviewResult.tenant_id == principal.tenant_id,
                ReviewResult.review_task_id == changed_request.review_task_id,
                ReviewResult.run_id == changed_request.run_id,
            )
        )
        if stale_result_count != 0:
            raise AssertionError("stale snapshot was published as a result")

    with factory() as session:
        set_tenant(session, "legal-b")
        visible = session.scalar(
            select(func.count(ReviewEngineRun.run_id)).where(
                ReviewEngineRun.tenant_id == principal.tenant_id,
                ReviewEngineRun.review_task_id == request.review_task_id,
            )
        )
        if visible != 0:
            raise AssertionError("review-run RLS exposed another tenant's review")
    assert runtime_settings.database_url is not None
    raw_url = runtime_settings.database_url.get_secret_value()
    raw_url = raw_url.replace("postgresql+psycopg://", "postgresql://", 1)
    thread_id = f"{principal.tenant_id}:{request.review_task_id}:{request.run_id}"
    with psycopg.connect(raw_url, autocommit=True, row_factory=dict_row) as conn:
        conn.execute("SELECT set_config('app.tenant_id', 'legal-b', false)")
        other_tenant = conn.execute(
            "SELECT count(*) AS n FROM checkpoints WHERE thread_id = %s", (thread_id,)
        ).fetchone()
        if other_tenant is None or other_tenant["n"] != 0:
            raise AssertionError("checkpoint RLS exposed another tenant's review state")
        conn.execute("SELECT set_config('app.tenant_id', 'legal-a', false)")
        own_tenant = conn.execute(
            "SELECT count(*) AS n FROM checkpoints WHERE thread_id = %s", (thread_id,)
        ).fetchone()
        if own_tenant is None or own_tenant["n"] == 0:
            raise AssertionError("checkpoint was not durably stored for its owning tenant")
    print(
        "PASS: real PostgreSQL checkpoint migration, offline gateway data flow, bounded retry, "
        "new-process interrupt recovery, reviewer decision, bounded retry, cancellation, "
        "revoked-identity and changed-snapshot rejection, RLS isolation, duplicate delivery, "
        "concurrent idempotent publication"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--resume", nargs=2, metavar=("TASK_ID", "RUN_ID"))
    args = parser.parse_args()
    if args.resume:
        _resume_worker(*args.resume)
    else:
        raise SystemExit("Invoke from check_legal_publications.py or pass --resume TASK_ID RUN_ID")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
