"""Create a clearly synthetic DOCX for the local offline interview demo."""

from __future__ import annotations

import argparse
import hashlib
from io import BytesIO
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / ".verification/task0020-demo/synthetic-tender.docx"


def create_fixture(output: Path) -> tuple[Path, str]:
    resolved = output.resolve()
    verification_root = (ROOT / ".verification").resolve()
    try:
        resolved.relative_to(verification_root)
    except ValueError as exc:
        raise ValueError("demo fixture output must stay under .verification") from exc
    resolved.parent.mkdir(parents=True, exist_ok=True)
    if resolved.exists():
        raise FileExistsError("refusing to overwrite an existing demo fixture")

    document = Document()
    document.core_properties.title = "招标文件合规审查演示数据（完全合成）"
    document.core_properties.subject = "TASK-0020 offline demo fixture"
    document.add_heading("招标文件（合成演示数据）", level=0)
    document.add_paragraph(
        "重要说明：本文件由程序生成的合成数据构成，不对应任何真实采购项目，"
        "不含真实招标文件或法规原文。"
    )
    document.add_heading("资格条件", level=1)
    document.add_paragraph("投标人必须在本市注册。该条款仅供离线演示和软件流程验证使用。")
    document.add_heading("其他说明", level=1)
    document.add_paragraph("请由审核员核对系统展示的原文定位、法规状态与证据充分性。")

    buffer = BytesIO()
    document.save(buffer)
    content = buffer.getvalue()
    with resolved.open("xb") as stream:
        stream.write(content)
    return resolved, hashlib.sha256(content).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    path, digest = create_fixture(args.output)
    print(f"Synthetic demo DOCX created: {path}")
    print(f"SHA-256: {digest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
