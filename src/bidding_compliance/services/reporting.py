"""Render structured immutable review results into Chinese Word and PDF reports."""

from __future__ import annotations

from io import BytesIO
from typing import Any


def _text(value: object, default: str = "未提供") -> str:
    return value if isinstance(value, str) and value.strip() else default


def _assessment_scope_lines(report: dict[str, Any]) -> list[str]:
    scope = report.get("assessment_scope")
    if not isinstance(scope, dict):
        return []
    lines: list[str] = []
    candidate_count = scope.get("candidate_count")
    if isinstance(candidate_count, int):
        if "model_candidate_count" in scope:
            lines.append(
                f"本次候选共 {candidate_count} 项；"
                f"确定性 Skills {scope.get('skill_candidate_count', 0)} 项，"
                f"AI 独立发现 {scope['model_candidate_count']} 项。"
            )
        else:
            lines.append(f"本次确定性 Skills 实际产生候选 {candidate_count} 项。")
        if candidate_count == 0:
            lines.append(
                "0 候选仅表示已运行规则范围内未产生候选，不代表标书合规、无风险或未发现违法。"
            )

    skill_names = {
        "qualification_review": "资格审查",
        "technical_review": "技术要求",
        "evaluation_review": "评审办法",
        "procedure_review": "采购程序",
        "contract_review": "合同条款",
    }
    status_names = {
        "completed_no_candidates": "已完成规则评估，未产生候选",
        "incomplete": "未完成判断",
        "not_applicable": "不适用",
        "completed": "已完成",
        "failed": "执行失败",
    }
    coverage = scope.get("skill_coverage", [])
    if isinstance(coverage, list) and coverage:
        lines.append(f"Skill 覆盖状态（{len(coverage)} 类）：")
        for item in coverage:
            if not isinstance(item, dict):
                continue
            skill_id = item.get("skill_id", "unknown")
            status = item.get("status", "unknown")
            name = skill_names.get(skill_id, str(skill_id))
            state = status_names.get(status, str(status))
            reason = item.get("reason")
            if isinstance(reason, str) and reason:
                reason_names = {
                    "ambiguous_score_table": "评分表存在歧义",
                    "ambiguous_tender_times": "采购时间存在歧义",
                }
                state += f"（{reason_names.get(reason, reason)}）"
            lines.append(f"{name}：{state}。")

    framework_confirmed = scope.get("legal_framework_confirmed")
    if framework_confirmed is False:
        lines.append("法律适用框架尚未确认。")
    elif framework_confirmed is True:
        lines.append("运行上下文记录了法律适用框架；本报告不据此扩大法规有效性结论。")

    errors = scope.get("retrieval_context_errors", {})
    status_counts = scope.get("retrieval_status_counts", {})
    hits = scope.get("retrieval_hit_count")
    citable = scope.get("citable_legal_evidence_count")
    if isinstance(errors, dict) and errors.get("CONTEXT_INVALID"):
        lines.append(
            "法规检索因 CONTEXT_INVALID 未能完成适用性匹配；未取得可确认、可引用的法规依据。"
        )
    elif isinstance(status_counts, dict) and status_counts.get("context_incomplete"):
        lines.append("法规检索上下文不完整，未完成有效适用性匹配。")
    if isinstance(hits, int) and isinstance(citable, int):
        lines.append(f"法规检索命中 {hits} 项，可引用依据 {citable} 项。")
    adoption = scope.get("legal_adoption_policy")
    if isinstance(adoption, dict):
        lines.append(
            f"法规采用政策：{adoption.get('policy_id')}（{adoption.get('decision_ref')}，"
            f"{adoption.get('adopted_on')}）；采用库内 {adoption.get('version_count')} 个版本/"
            f"{adoption.get('fragment_count')} 个片段；未做独立法律核验。"
        )
        lines.append(
            f"项目采用参考 {scope.get('adopted_legal_reference_count', 0)} 条；"
            f"其中工程配置规则—法规关联 {scope.get('configured_rule_law_binding_count', 0)} 条；"
            f"正式可引用（独立核验）依据 {citable} 条。"
        )

    provider = scope.get("model_provider")
    model_name = scope.get("model_name")
    if provider == "offline-verification":
        lines.append(
            f"语义分析使用 OfflineReviewGateway（{provider} / {model_name}）；本次未调用真实 LLM。"
        )
    elif isinstance(provider, str):
        lines.append(f"语义分析提供方：{provider}；模型：{_text(model_name)}。")
    if scope.get("model_execution_mode") == "real-qwen":
        lines.append(
            f"真实模型已验证调用 {scope.get('model_calls', 0)} 次；"
            f"输入/输出 tokens {scope.get('model_input_tokens', 0)}/"
            f"{scope.get('model_output_tokens', 0)}；"
            f"按官方价格估算 ¥{scope.get('model_cost_cny', 0):.6f}。"
        )
    indexed, total = (
        scope.get("vector_index_item_count", 0),
        scope.get("vector_source_item_count", 0),
    )
    if total:
        lines.append(
            f"真实神经向量索引覆盖法规片段 {indexed}/{total}；其余保留 BM25。相似度不是法律证据。"
        )

    product = scope.get("product_execution")
    if isinstance(product, dict):
        completion = product.get("completion_status")
        coverage_status = product.get("coverage_status")
        if completion == "partial":
            lines.append(
                "本次审查为部分完成：存在未提交、调用失败、覆盖无法确认或结果未知的"
                "section，不得解释为全部内容已审查完成。"
            )
        elif completion == "complete" and coverage_status == "complete":
            lines.append("本次审查覆盖完整：所有计划单元均已处理，未发现饱和或结果未知的调用。")
        elif coverage_status == "not_applicable":
            lines.append("本次审查未使用模型发现覆盖计划；覆盖状态不适用于确定性规则路径。")
        saturated_units = product.get("saturated_units")
        if isinstance(saturated_units, int) and saturated_units > 0:
            lines.append(
                f"已提交模型但达到单批发现上限、覆盖无法确认的 section：{saturated_units} 个；"
                "这不是完全未处理，但报告不声称这些 section 已覆盖全部疑点。"
            )
        rejected_units = product.get("rejected_units")
        if isinstance(rejected_units, int) and rejected_units > 0:
            lines.append(f"已分析但候选被拒绝的 section：{rejected_units} 个。")
        unprocessed = product.get("unprocessed_units")
        if isinstance(unprocessed, int) and unprocessed > 0:
            lines.append(
                f"未提交模型或调用失败的 section：{unprocessed} 个，需人工处理。"
            )
        truncated_batches = product.get("truncated_batches")
        if isinstance(truncated_batches, int) and truncated_batches > 0:
            lines.append(f"模型响应截断的批次：{truncated_batches} 个，覆盖不完整。")
        elif product.get("saturated") and not saturated_units:
            lines.append("模型发现存在饱和批次；覆盖状态按部分完成处理。")
        unknown_calls = product.get("unknown_model_call_count")
        if isinstance(unknown_calls, int) and unknown_calls > 0:
            lines.append(
                f"结果未知的模型调用 {unknown_calls} 次；未自动重放，费用按预留上限保留。"
            )

    confirmation = scope.get("no_findings_confirmation")
    if confirmation == "confirm_no_findings":
        lines.append(
            "confirm_no_findings 属于演示流程的人工作业记录，不构成专业法律审查确认，"
            "也不表示全部合规。"
        )
    elif confirmation is None and candidate_count == 0:
        lines.append("未找到零候选人工确认记录；不得将本结果解释为专业法律审查结论。")
    return lines


def render_docx(report: dict[str, Any]) -> bytes:
    from docx import Document as WordDocument
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = WordDocument()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.different_first_page_header_footer = False
    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10 if style_name == "Normal" else 16 if style_name == "Title" else 12)
    header = section.header.paragraphs[0]
    header.text = "招标文件智能合规审查报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"报告 {report['report_id']}  |  结果 {report['result_id']}  |  第 ")
    _word_field(footer, "PAGE")
    footer.add_run(" 页")

    title = document.add_heading("招标文件合规审查报告", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (
        ("项目", _text(report.get("project_name"), report["project_id"])),
        ("审查任务", report["review_task_id"]),
        ("结果版本", report["result_id"]),
        ("执行批次", report["run_id"]),
        ("输入快照", report["snapshot_hash"]),
        ("报告版本", report["report_id"]),
        ("生成时间（UTC）", report["generated_at"]),
    ):
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = label, str(value)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    document.add_heading("一、审查范围与说明", level=1)
    document.add_paragraph(
        "本报告依据已发布的不可变审查结果快照生成。报告记录候选风险、可核验的证据状态及人工复核决定。"
        "系统当前未批准规则与法规条款的充分支持绑定，证据状态最高为 suspected；"
        "本报告不构成法律意见，"
        "不替代采购人及专业人员的审查责任。"
    )
    snapshot = report.get("input_snapshot", {})
    document.add_paragraph("被审查文档版本：")
    docs = snapshot.get("documents", [])
    if docs:
        for item in docs:
            document.add_paragraph(
                f"文档 {item['document_id']}，SHA-256 {item['raw_source_sha256']}，"
                f"解析版本 {item['parse_version']}，上下文版本 {item['context_version']}。",
                style="List Bullet",
            )
    else:
        document.add_paragraph("快照没有文档清单。")

    findings = report["findings"]
    document.add_heading("二、风险汇总", level=1)
    rejected = sum(item["review_status"] == "rejected" for item in findings)
    accepted = sum(item["review_status"] == "accepted" for item in findings)
    modified = sum(item["review_status"] == "modified" for item in findings)
    document.add_paragraph(
        f"风险条目 {len(findings)} 项；接受 {accepted} 项，修改 {modified} 项，驳回 {rejected} 项。"
        "驳回项目不计入已接受风险。"
    )
    summary = document.add_table(rows=1, cols=4)
    summary.style = "Table Grid"
    for cell, label in zip(summary.rows[0].cells, ("类别", "条目", "接受", "驳回"), strict=True):
        cell.text = label
    categories = ("qualification", "technical", "evaluation", "contract", "procedure")
    category_names = {
        "qualification": "资格审查",
        "technical": "技术要求",
        "evaluation": "评标办法",
        "contract": "合同条款",
        "procedure": "招标程序",
    }
    for category in categories:
        rows = [item for item in findings if item["category"] == category]
        row = summary.add_row()
        for cell, value in zip(
            row.cells,
            (
                category_names[category],
                str(len(rows)),
                str(sum(item["review_status"] == "accepted" for item in rows)),
                str(sum(item["review_status"] == "rejected" for item in rows)),
            ),
            strict=True,
        ):
            cell.text = value

    document.add_heading("三、风险及复核明细", level=1)
    if not findings:
        document.add_paragraph("本次审查未发现候选风险；该表述仅对应此任务与输入快照。")
    for index, finding in enumerate(findings, start=1):
        document.add_heading(f"{index}. {finding['title']}", level=2)
        document.add_paragraph(
            f"类别：{category_names[finding['category']]}　严重程度：{finding['severity']}"
        )
        document.add_paragraph(f"描述：{finding['description']}")
        document.add_paragraph(
            f"证据状态：{finding['evidence_status']}（独立 Evidence Checker 结果）"
        )
        document.add_paragraph(f"人工复核：{finding['review_status']}")
        review = finding.get("human_review", {})
        if review:
            document.add_paragraph(
                f"复核人：{review['reviewer_id']}　时间：{review['decided_at']}\n"
                f"理由：{review['reason']}"
            )
        document.add_paragraph(f"调整建议：{finding['recommendation']}")
        for cite in finding["citations"]:
            location = cite.get("page")
            page_label = f"第 {location} 页" if location is not None else "原文无可靠页码"
            document.add_paragraph(
                f"原文证据：文档 {cite['document_id']}，版本 {cite['raw_source_sha256']}，"
                f"{page_label}，定位 {cite['locator']}，"
                f"片段 {cite['section_id']}@{cite['section_revision']}。"
            )
            document.add_paragraph(f"“{cite['quote']}”")
        laws = [
            item
            for item in finding.get("evidence", [])
            if item.get("kind") in {"law", "policy", "internal_rule"}
        ]
        if laws:
            for law in laws:
                reference = f"{law.get('title', '法规证据')} {law.get('article') or '条款未提供'}"
                document.add_paragraph(
                    f"法规或规则证据：{reference}；状态 {law['version_status']}。"
                )
                document.add_paragraph(f"“{law['quote']}”")
        else:
            document.add_paragraph("法规或规则证据：本条结果未包含可引用的法规条款，不补造条文。")
        references = finding.get("legal_references", [])
        if references:
            adoption = report.get("assessment_scope", {}).get("legal_adoption_policy")
            document.add_paragraph(
                "项目采用法规库参考（未经独立法律核验，不计入已核验证据）："
                if isinstance(adoption, dict)
                else "可能相关依据（适用性未经确认，仅供参考，不计入已核验证据）："
            )
            for reference_item in references[:3]:
                label = (
                    f"{reference_item.get('title', '')} "
                    f"{reference_item.get('article') or '条款未提供'}"
                )
                source_kind = (
                    "模型引用候选"
                    if reference_item.get("selected_by_model")
                    else "检索候选"
                )
                suffix = (
                    "；项目采用库内版本（未经独立法律核验）"
                    if reference_item.get("adoption_status") == "project_adopted"
                    else ""
                )
                binding = reference_item.get("rule_law_binding_id")
                if binding:
                    suffix += f"；工程配置规则—法规关联 {binding}"
                document.add_paragraph(
                    f"参考法规（{source_kind}）：{label}；版本状态 "
                    f"{reference_item.get('version_status')}；未确认原因 "
                    f"{reference_item.get('confirmation_reason')}{suffix}。"
                )
                document.add_paragraph(f"“{reference_item.get('quote', '')}”")
    scope_lines = _assessment_scope_lines(report)
    if scope_lines:
        document.add_heading("四、本次运行范围与限制", level=1)
        for line in scope_lines:
            document.add_paragraph(line, style="List Bullet")
    document.add_heading("五、限制与免责声明" if scope_lines else "四、限制与免责声明", level=1)
    document.add_paragraph(
        "报告只呈现指定不可变结果版本。证据不足或规则关联未批准时，不得据此认定确定性法律违规；"
        "原文页码仅在解析数据提供且已校验时显示。报告不含伪造签名或审批效力。"
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _word_field(paragraph: Any, name: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = name
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def render_pdf(report: dict[str, Any]) -> bytes:
    from reportlab.lib import colors  # type: ignore[import-untyped]
    from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore[import-untyped]
    from reportlab.lib.units import mm  # type: ignore[import-untyped]
    from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
    from reportlab.platypus import (  # type: ignore[import-untyped]
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "CJKTitle",
            parent=styles["Title"],
            fontName="STSong-Light",
            fontSize=18,
            leading=24,
            alignment=1,
        )
    )
    styles.add(
        ParagraphStyle(
            "CJKHeading",
            parent=styles["Heading2"],
            fontName="STSong-Light",
            fontSize=12,
            leading=18,
            spaceBefore=8,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            "CJKBody",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=9,
            leading=14,
            wordWrap="CJK",
        )
    )

    def para(value: object, style: str = "CJKBody") -> Paragraph:
        from xml.sax.saxutils import escape

        return Paragraph(escape(str(value)).replace("\n", "<br/>"), styles[style])

    story: list[Any] = [para("招标文件合规审查报告", "CJKTitle"), Spacer(1, 8)]
    meta = [
        [para("项目"), para(_text(report.get("project_name"), report["project_id"]))],
        [para("审查任务 / 结果版本"), para(f"{report['review_task_id']} / {report['result_id']}")],
        [para("运行 / 输入快照"), para(f"{report['run_id']} / {report['snapshot_hash']}")],
        [
            para("报告版本 / 生成时间 UTC"),
            para(f"{report['report_id']} / {report['generated_at']}"),
        ],
    ]
    table = Table(meta, colWidths=(42 * mm, 128 * mm), repeatRows=0)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend([table, Spacer(1, 10), para("一、审查说明", "CJKHeading")])
    story.append(
        para(
            "本报告依据已发布的不可变结果快照生成。规则与法规条款的充分支持绑定尚未批准，"
            "证据状态最高为 suspected。本报告不构成法律意见。"
        )
    )
    story.append(para("二、风险汇总", "CJKHeading"))
    findings = report["findings"]
    counts = [
        sum(item["review_status"] == status for item in findings)
        for status in ("accepted", "modified", "rejected")
    ]
    story.append(
        para(
            f"共 {len(findings)} 项；接受 {counts[0]} 项，修改 {counts[1]} 项，"
            f"驳回 {counts[2]} 项。"
            "驳回项目不计入已接受风险。"
        )
    )
    labels = {
        "qualification": "资格审查",
        "technical": "技术要求",
        "evaluation": "评标办法",
        "contract": "合同条款",
        "procedure": "招标程序",
    }
    summary_rows = [[para("类别"), para("条目"), para("接受"), para("驳回")]]
    for category, name in labels.items():
        rows = [item for item in findings if item["category"] == category]
        summary_rows.append(
            [
                para(name),
                para(len(rows)),
                para(sum(x["review_status"] == "accepted" for x in rows)),
                para(sum(x["review_status"] == "rejected" for x in rows)),
            ]
        )
    summary = Table(summary_rows, colWidths=(90 * mm, 30 * mm, 25 * mm, 25 * mm), repeatRows=1)
    summary.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.extend([summary, Spacer(1, 10), para("三、风险及复核明细", "CJKHeading")])
    if not findings:
        story.append(para("本次审查未发现候选风险；该表述仅对应此任务与输入快照。"))
    for index, item in enumerate(findings, 1):
        story.append(para(f"{index}. {item['title']}", "CJKHeading"))
        story.append(
            para(
                f"类别：{labels[item['category']]}；严重程度：{item['severity']}；证据状态：{item['evidence_status']}；人工复核：{item['review_status']}"
            )
        )
        story.append(para(f"描述：{item['description']}"))
        review = item.get("human_review", {})
        if review:
            story.append(
                para(
                    f"复核人：{review['reviewer_id']}；时间：{review['decided_at']}；理由：{review['reason']}"
                )
            )
        story.append(para(f"调整建议：{item['recommendation']}"))
        for cite in item["citations"]:
            page = f"第 {cite['page']} 页" if cite.get("page") is not None else "原文无可靠页码"
            story.append(
                para(
                    f"原文定位：文档 {cite['document_id']}，SHA-256 {cite['raw_source_sha256']}，"
                    f"{page}，{cite['locator']}，片段 "
                    f"{cite['section_id']}@{cite['section_revision']}。"
                )
            )
            story.append(para(f"原文摘录：{cite['quote']}"))
        laws = [
            e
            for e in item.get("evidence", [])
            if e.get("kind") in {"law", "policy", "internal_rule"}
        ]
        if laws:
            for law in laws:
                story.append(
                    para(
                        f"法规/规则：{law.get('title', '证据')} "
                        f"{law.get('article') or '条款未提供'}；状态 {law['version_status']}。"
                    )
                )
                story.append(para(law["quote"]))
        else:
            story.append(para("法规或规则证据：未提供可引用条款，不补造条文。"))
        references = item.get("legal_references", [])
        if references:
            adoption = report.get("assessment_scope", {}).get("legal_adoption_policy")
            story.append(
                para(
                    "项目采用法规库参考（未经独立法律核验，不计入已核验证据）："
                    if isinstance(adoption, dict)
                    else "可能相关依据（适用性未经确认，仅供参考，不计入已核验证据）："
                )
            )
            for reference_item in references[:3]:
                source_kind = (
                    "模型引用候选"
                    if reference_item.get("selected_by_model")
                    else "检索候选"
                )
                suffix = (
                    "；项目采用库内版本（未经独立法律核验）"
                    if reference_item.get("adoption_status") == "project_adopted"
                    else ""
                )
                binding = reference_item.get("rule_law_binding_id")
                if binding:
                    suffix += f"；工程配置规则—法规关联 {binding}"
                story.append(
                    para(
                        f"参考法规（{source_kind}）：{reference_item.get('title', '')} "
                        f"{reference_item.get('article') or '条款未提供'}；版本状态 "
                        f"{reference_item.get('version_status')}；未确认原因 "
                        f"{reference_item.get('confirmation_reason')}{suffix}。"
                    )
                )
                story.append(para(f"“{reference_item.get('quote', '')}”"))
    scope_lines = _assessment_scope_lines(report)
    if scope_lines:
        story.append(para("四、本次运行范围与限制", "CJKHeading"))
        story.extend(para(line) for line in scope_lines)
    story.extend(
        [
            para("五、限制与免责声明" if scope_lines else "四、限制与免责声明", "CJKHeading"),
            para(
                "本报告仅呈现指定结果版本。证据不足或规则关联未批准时，不得据此认定确定性法律违规。报告不含伪造签名或审批效力。"
            ),
        ]
    )
    buffer = BytesIO()

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            f"报告 {report['report_id']}  |  结果 {report['result_id']}  |  {document.page}",
        )
        canvas.restoreState()

    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="招标文件合规审查报告",
        author="招标文件智能合规审查系统",
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
