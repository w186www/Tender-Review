"""TASK0005-001/002: malicious inputs, durable processing, locations and human correction."""

import hashlib
import io
import subprocess
import sys
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from defusedxml.common import DefusedXmlException
from docx import Document as WordDocument
from PIL import Image
from pypdf import PdfWriter
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_identity_projects import create, sign_in
from test_identity_projects import system as shared_system
from verify_contracts import ROOT, load_contracts, load_openapi, validator_for_contract

from bidding_compliance.ingestion.parser import parse
from bidding_compliance.ingestion.policy import load_policy
from bidding_compliance.ingestion.ports import IngestionError
from bidding_compliance.ingestion.quality import digest, quality
from bidding_compliance.ingestion.schemas import Fragment
from bidding_compliance.ingestion.storage import LocalObjectStore
from bidding_compliance.persistence.models import (
    Document,
    DocumentContextVersion,
    DocumentSection,
    IngestionJob,
    User,
)
from bidding_compliance.services.ingestion_worker import execute

system = shared_system


def pdf_payload(text="Synthetic tender qualification and contract source text."):
    stream = io.BytesIO()
    document = canvas.Canvas(stream)
    document.drawString(50, 700, text)
    document.showPage()
    document.save()
    return stream.getvalue()


def docx_payload():
    stream = io.BytesIO()
    document = WordDocument()
    document.add_heading("Synthetic qualification section", 1)
    document.add_paragraph("政府采购工程。项目名称：合成学校改造。原始资格要求仅用于定位测试。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Requirement"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Synthetic capacity"
    table.cell(1, 1).text = "Test only"
    document.save(stream)
    return stream.getvalue()


class UnitScanner:
    """Explicit unit double; actual ClamAV is exercised by the mandatory container gate."""

    def __init__(self, error=None):
        self.error = error

    def scan(self, payload):
        if self.error:
            raise IngestionError(self.error)


class UnitParser:
    def __init__(self, error=None):
        self.error = error
        self.calls = 0

    def parse(self, payload, format, policy):
        self.calls += 1
        if self.error:
            raise IngestionError(self.error)
        return parse(payload, format, policy, ocr=UnitOCR())


class UnitOCR:
    def recognize(self, image):
        return []


def upload(system, payload=None, filename="synthetic.pdf", media="application/pdf", headers=None):
    client, _, _, _ = system
    headers = sign_in(system) if headers is None else headers
    project = create(client, headers).json()["project_id"]
    response = client.post(
        f"/api/v1/projects/{project}/documents",
        headers={**headers, "Idempotency-Key": str(uuid4())},
        files={"file": (filename, payload if payload is not None else pdf_payload(), media)},
    )
    return project, response


def process(system, document, parser=None, scanner=None):
    _, engine, settings, _ = system
    return execute(
        engine,
        "a",
        document,
        parser or UnitParser(),
        scanner or UnitScanner(),
        LocalObjectStore(settings.object_storage_root),
        load_policy(),
    )


def validate(name, data):
    documents, registry, errors = load_contracts(ROOT)
    api, api_errors = load_openapi(ROOT)
    assert not errors + api_errors
    validator_for_contract(
        ROOT, "contracts/openapi.yaml#/components/schemas/" + name, documents, registry, api
    ).validate(data)


def test_parser_contract_generator_has_no_drift():
    result = subprocess.run(  # noqa: S603 - fixed repository generator
        [sys.executable, "scripts/generate_ingestion_contracts.py", "--check"],
        cwd=ROOT,
        capture_output=True,
        check=False,
        timeout=10,
    )
    assert result.returncode == 0, result.stdout


@pytest.mark.parametrize(
    "filename,media,payload,status",
    [
        ("bad.exe", "application/pdf", b"%PDF-1.7", 415),
        ("bad.pdf", "image/png", pdf_payload(), 415),
        ("bad.pdf", "application/pdf", b"not a PDF", 415),
        ("../bad.pdf", "application/pdf", pdf_payload(), 415),
        ("empty.pdf", "application/pdf", b"", 413),
        ("large.pdf", "application/pdf", b"%PDF-" + b"x" * 10485761, 413),
    ],
    ids=["extension", "mime", "signature", "path", "empty", "size"],
)
def test_invalid_upload_rejected_before_job_or_object(system, filename, media, payload, status):
    _, response = upload(system, payload, filename, media)
    assert response.status_code == status, response.text
    validate("Error", response.json())
    _, engine, settings, _ = system
    with Session(engine) as db:
        assert list(db.scalars(select(IngestionJob))) == []
    assert not list(Path(settings.object_storage_root).rglob("*"))


def test_native_pdf_locations_and_closed_parsed_contract():
    payload = pdf_payload()
    parsed = parse(payload, "pdf", load_policy(), UnitOCR())
    assert parsed.quality == "passed" and parsed.page_count == 1
    assert parsed.source_hash == hashlib.sha256(payload).hexdigest()
    first = parsed.fragments[0]
    assert first.page == 1 and first.bbox and first.coordinate_system == "pdf_points"
    assert first.text == "Synthetic tender qualification and contract source text."
    assert first.text_sha256 == hashlib.sha256(first.text.encode()).hexdigest()
    assert first.locator.startswith("pdf/page/1/")


def test_docx_preserves_heading_table_cells_and_does_not_invent_pages():
    parsed = parse(docx_payload(), "docx", load_policy(), UnitOCR())
    assert parsed.quality == "passed"
    assert any(f.kind == "table_cell" and f.row == 1 and f.column == 0 for f in parsed.fragments)
    assert all(f.page is None and f.bbox is None for f in parsed.fragments)
    assert any(f.section == "Synthetic qualification section" for f in parsed.fragments)
    assert len({f.locator for f in parsed.fragments}) == len(parsed.fragments)


@pytest.mark.parametrize("kind", ["javascript", "encrypted", "pages"])
def test_unsafe_pdf_rejected_inside_parser(kind):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    expected = "MALICIOUS_CONTENT"
    policy = load_policy()
    if kind == "javascript":
        writer.add_js("app.alert('synthetic');")
    elif kind == "encrypted":
        writer.encrypt(str(uuid4()))
        expected = "ENCRYPTED_FILE"
    else:
        writer.add_blank_page(width=100, height=100)
        policy = policy.model_copy(update={"max_pages": 1})
        expected = "PAGE_LIMIT"
    stream = io.BytesIO()
    writer.write(stream)
    with pytest.raises(IngestionError) as error:
        parse(stream.getvalue(), "pdf", policy, UnitOCR())
    assert error.value.code == expected


@pytest.mark.parametrize("kind", ["traversal", "macro", "external", "bomb", "doctype"])
def test_malicious_docx_rejected_inside_parser(kind):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as output:
        with zipfile.ZipFile(io.BytesIO(docx_payload())) as source:
            for name in source.namelist():
                data = source.read(name)
                if kind == "external" and name == "_rels/.rels":
                    data = b'<Relationships><Relationship TargetMode="External" Target="http://invalid"/></Relationships>'
                if kind == "doctype" and name == "word/document.xml":
                    data = b'<!DOCTYPE x [<!ENTITY a "synthetic">]><x>&a;</x>'
                output.writestr(name, data)
        if kind == "traversal":
            output.writestr("../escaped", "synthetic")
        elif kind == "macro":
            output.writestr("word/vbaProject.bin", "synthetic")
        elif kind == "bomb":
            output.writestr("word/bomb.xml", "x" * 1000000)
    with pytest.raises((IngestionError, DefusedXmlException)):
        parse(stream.getvalue(), "docx", load_policy(), UnitOCR())


def test_document_lifecycle_and_context_version_history(system):
    client, engine, _, _ = system
    payload = docx_payload()
    project, uploaded = upload(
        system,
        payload,
        "synthetic.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert uploaded.status_code == 202, uploaded.text
    validate("Document", uploaded.json())
    document = uploaded.json()["document_id"]
    path = f"/api/v1/projects/{project}/documents/{document}"
    headers = sign_in(system)
    assert client.get(path + "/original", headers=headers).status_code == 409
    assert process(system, document).status == "needs_confirmation"
    section_response = client.get(path + "/sections", headers=headers)
    assert section_response.status_code == 200, section_response.text
    validate("DocumentSections", section_response.json())
    context = client.get(path + "/context", headers=headers).json()
    validate("DocumentContext", context)
    assert any(v["field"] == "procurement_type" for v in context["candidates"])
    decision = {
        "expected_version": 1,
        "reason": "合成人工确认",
        "context": dict(
            name="synthetic",
            procurement_type="government",
            project_type="works",
            region="synthetic",
        ),
    }
    confirmed = client.post(path + "/context", headers=headers, json=decision)
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "ready"
    assert confirmed.json()["context_version"] == 2
    assert client.post(path + "/context", headers=headers, json=decision).status_code == 409
    original = client.get(path + "/original", headers=headers)
    assert original.status_code == 200 and original.content == payload
    with Session(engine) as db:
        snapshots = list(
            db.scalars(select(DocumentContextVersion).order_by(DocumentContextVersion.version))
        )
        assert [s.version for s in snapshots] == [1, 2]
        assert snapshots[0].payload["manual"] is None
        assert snapshots[1].payload["reason"] == "合成人工确认"


def test_low_quality_manual_correction_rechecks_quality_and_preserves_automatic(system):
    client, _, _, _ = system
    project, uploaded = upload(system, pdf_payload("Short"))
    document = uploaded.json()["document_id"]
    assert process(system, document).status == "needs_review"
    headers = sign_in(system)
    path = f"/api/v1/projects/{project}/documents/{document}"
    confirmation = {
        "expected_version": 1,
        "reason": "test",
        "context": {
            "name": "test",
            "procurement_type": "government",
            "project_type": "works",
            "region": "test",
        },
    }
    assert client.post(path + "/context", headers=headers, json=confirmation).status_code == 409
    items = client.get(path + "/sections", headers=headers).json()["items"]
    # Fix every OCR placeholder as an explicitly human transcript; do not change automatic evidence.
    correction = {
        "expected_parse_version": 1,
        "reason": "合成人工转录",
        "corrections": [
            {
                "section_id": item["section_id"],
                "text": "Synthetic manually verified source transcription.",
            }
            for item in items
        ],
    }
    fixed = client.post(path + "/corrections", headers=headers, json=correction)
    assert fixed.status_code == 200, fixed.text
    assert fixed.json()["status"] == "needs_confirmation"
    assert fixed.json()["parse_version"] == 2
    changed = client.get(path + "/sections", headers=headers).json()["items"]
    assert changed[0]["automatic"]["text"] == items[0]["automatic"]["text"]
    assert all(item["effective"]["source"] == "manual" for item in changed)
    assert client.post(path + "/corrections", headers=headers, json=correction).status_code == 409


@pytest.mark.parametrize("code", ["MALICIOUS_CONTENT", "SCANNER_UNAVAILABLE", "SCANNER_STALE"])
def test_scan_failure_cannot_parse_or_be_manually_overridden(system, code):
    client, _, _, _ = system
    project, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    parser = UnitParser()
    result = process(system, document, parser, UnitScanner(code))
    assert result.status in {"rejected", "failed"} and result.issues == [code]
    assert parser.calls == 0
    path = f"/api/v1/projects/{project}/documents/{document}"
    headers = sign_in(system)
    assert client.get(path + "/original", headers=headers).status_code == 409
    assert client.get(path + "/sections", headers=headers).status_code == 409


def test_cross_tenant_access_upload_and_role_writes_are_rejected(system):
    client, _, _, _ = system
    project, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    assert process(system, document).status == "needs_confirmation"
    path = f"/api/v1/projects/{project}/documents/{document}"
    other = sign_in(system, tenant="b")
    for suffix in ("", "/original", "/sections", "/context"):
        assert client.get(path + suffix, headers=other).status_code == 404
    for role in ("admin", "operator", "auditor"):
        headers = sign_in(system, role=role)
        rejected = client.post(
            f"/api/v1/projects/{project}/documents",
            headers={**headers, "Idempotency-Key": str(uuid4())},
            files={"file": ("synthetic.pdf", pdf_payload(), "application/pdf")},
        )
        assert rejected.status_code == 403


def test_worker_rechecks_submitter_and_hash_and_retry_budget(system):
    _, engine, settings, _ = system
    _, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    with Session(engine) as db:
        user = db.get(User, ("a", "reviewer"))
        user.active = False
        db.commit()
    assert process(system, document).issues == ["ACTOR_FORBIDDEN"]
    with Session(engine) as db:
        user = db.get(User, ("a", "reviewer"))
        user.active = True
        db.commit()
    _, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    with Session(engine) as db:
        row = db.get(Document, ("a", document))
        path = LocalObjectStore(settings.object_storage_root).path(row.object_ref)
    path.write_bytes(b"modified synthetic payload")
    assert process(system, document).issues == ["HASH_MISMATCH"]
    _, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    for _ in range(3):
        assert (
            process(system, document, scanner=UnitScanner("SCANNER_UNAVAILABLE")).status == "failed"
        )
    with pytest.raises(FileExistsError):
        process(system, document)


def test_empty_ocr_page_and_broken_font_mapping_cannot_be_hidden_by_other_text():
    native = Fragment(
        locator="ooxml/body/0",
        coordinate_system="ooxml",
        kind="paragraph",
        text="Long valid text covering a different source page.",
        text_sha256=digest("Long valid text covering a different source page."),
        source="native",
    )
    blank = Fragment(
        locator="pdf/page/2/ocr/empty",
        page=2,
        bbox=[0, 0, 100, 100],
        coordinate_system="pdf_points",
        kind="ocr",
        text="",
        text_sha256=digest(""),
        source="ocr",
        confidence=0,
    )
    assert not quality([native, blank], load_policy())
    broken = native.model_copy(update={"text": "(cid:12)" * 20})
    assert not quality([broken], load_policy())


def test_docx_footer_nested_table_and_embedded_image_do_not_disappear():
    document = WordDocument()
    document.add_paragraph("Synthetic body source text with qualification requirements.")
    document.sections[0].footer.paragraphs[0].text = "Synthetic footer qualification condition."
    parent = document.add_table(rows=1, cols=1)
    nested = parent.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Synthetic nested table condition."
    picture = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(picture, format="PNG")
    picture.seek(0)
    document.add_picture(picture)
    output = io.BytesIO()
    document.save(output)
    parsed = parse(output.getvalue(), "docx", load_policy(), UnitOCR())
    assert any("footer qualification" in f.text for f in parsed.fragments)
    assert any("nested table condition" in f.text for f in parsed.fragments)
    assert any(f.locator.startswith("ooxml/embedded-image/") for f in parsed.fragments)
    assert parsed.quality == "requires_review" and parsed.issues == ["LOW_QUALITY"]


def test_raw_upload_body_limit_keeps_correlated_error(system):
    client, _, _, _ = system
    response = client.post(
        "/api/v1/projects/synthetic/documents",
        content=b"x",
        headers={"Content-Length": str(load_policy().max_upload_bytes + 65537)},
    )
    assert response.status_code == 413
    validate("Error", response.json())
    assert response.headers["x-request-id"] == response.json()["request_id"]


def test_worker_lease_prevents_duplicate_and_superseded_publication(system):
    _, engine, _, _ = system
    _, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    with Session(engine) as db:
        job = db.get(IngestionJob, ("a", document))
        job.status, job.lease_id = "running", str(uuid4())
        job.lease_until = datetime.now(UTC) + timedelta(minutes=3)
        db.commit()
    parser = UnitParser()
    with pytest.raises(FileExistsError):
        process(system, document, parser)
    assert parser.calls == 0
    with Session(engine) as db:
        job = db.get(IngestionJob, ("a", document))
        job.lease_until = datetime.now(UTC) - timedelta(seconds=1)
        db.commit()

    class TakeoverScanner:
        def scan(self, payload):
            with Session(engine) as db:
                job = db.get(IngestionJob, ("a", document))
                job.lease_id = str(uuid4())
                db.commit()

    with pytest.raises(FileExistsError):
        process(system, document, scanner=TakeoverScanner())
    with Session(engine) as db:
        assert list(db.scalars(select(DocumentSection))) == []
        assert list(db.scalars(select(DocumentContextVersion))) == []


def test_submitter_revoked_during_parse_cannot_publish(system):
    _, engine, _, _ = system
    _, uploaded = upload(system)
    document = uploaded.json()["document_id"]

    class RevocationScanner:
        def scan(self, payload):
            with Session(engine) as db:
                db.get(User, ("a", "reviewer")).active = False
                db.commit()

    assert process(system, document, scanner=RevocationScanner()).issues == ["ACTOR_FORBIDDEN"]
    with Session(engine) as db:
        assert list(db.scalars(select(DocumentSection))) == []


def test_text_correction_racing_context_confirmation_returns_conflict(system, monkeypatch):
    from bidding_compliance.services import documents

    client, _, _, _ = system
    project, uploaded = upload(system)
    document = uploaded.json()["document_id"]
    assert process(system, document).status == "needs_confirmation"
    headers = sign_in(system)
    path = f"/api/v1/projects/{project}/documents/{document}"
    items = client.get(path + "/sections", headers=headers).json()["items"]
    original_quality = documents.quality

    def confirm_during_correction(fragments, policy):
        confirmation = client.post(
            path + "/context",
            headers=headers,
            json={
                "expected_version": 1,
                "reason": "Synthetic concurrent confirmation",
                "context": dict(
                    name="synthetic",
                    procurement_type="government",
                    project_type="works",
                    region="synthetic",
                ),
            },
        )
        assert confirmation.status_code == 200
        return original_quality(fragments, policy)

    monkeypatch.setattr(documents, "quality", confirm_during_correction)
    correction = client.post(
        path + "/corrections",
        headers=headers,
        json={
            "expected_parse_version": 1,
            "reason": "Synthetic concurrent transcription",
            "corrections": [
                {"section_id": items[0]["section_id"], "text": "Verified source text."}
            ],
        },
    )
    assert correction.status_code == 409, correction.text
    current = client.get(path, headers=headers).json()
    assert current["context_version"] == 2 and current["parse_version"] == 1
    assert client.get(path + "/context", headers=headers).json()["confirmed"]


def test_parser_response_exclusivity_and_table_indices_are_contract_constraints():
    from jsonschema import ValidationError

    parsed = parse(pdf_payload(), "pdf", load_policy(), UnitOCR()).model_dump()
    documents, registry, errors = load_contracts(ROOT)
    assert not errors
    validator = validator_for_contract(
        ROOT, "contracts/schemas/parser-response.schema.json", documents, registry, None
    )
    validator.validate({"result": parsed, "error": None})
    with pytest.raises(ValidationError):
        validator.validate({"result": parsed, "error": "INVALID_FILE"})
    with pytest.raises(ValidationError):
        validator.validate({"result": None, "error": None})
    parsed["fragments"][0]["kind"] = "table_cell"
    with pytest.raises(ValidationError):
        validator.validate({"result": parsed, "error": None})
