"""TASK0010 report rendering must preserve published evidence and Chinese text."""

import json
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from jsonschema import Draft202012Validator
from pypdf import PdfReader
from verify_contracts import load_contracts, load_openapi, validator_for_contract

from bidding_compliance.domain.identity import Principal, Role
from bidding_compliance.services.reporting import render_docx, render_pdf
from bidding_compliance.services.review_persistence import _decision_projection
from bidding_compliance.services.review_projection import project_product_execution


def _report() -> dict[str, object]:
    return {
        "schema_version": "task0010-report.v1",
        "tenant_id": "tenant-test",
        "project_id": "project-test",
        "project_name": "中文项目",
        "review_task_id": "task-test",
        "run_id": "run-test",
        "result_id": "result-test",
        "snapshot_hash": "a" * 64,
        "report_id": "report-test",
        "template_version": "task0010-v1",
        "generated_at": "2026-10-14T00:00:00+00:00",
        "input_snapshot": {
            "documents": [
                {
                    "document_id": "doc-old-version",
                    "raw_source_sha256": "b" * 64,
                    "parse_version": 1,
                    "context_version": 1,
                }
            ]
        },
        "findings": [
            {
                "finding_id": "finding-test",
                "category": "qualification",
                "title": "注册地要求需复核",
                "description": "招标文件要求供应商在本市注册。",
                "severity": "medium",
                "evidence_status": "suspected",
                "review_status": "accepted",
                "recommendation": "核实要求与项目需求的关联。",
                "human_review": {
                    "reviewer_id": "reviewer-test",
                    "decided_at": "2026-10-14T00:00:00+00:00",
                    "reason": "请结合项目实际复核。",
                },
                "evidence": [],
                "citations": [
                    {
                        "document_id": "doc-old-version",
                        "raw_source_sha256": "b" * 64,
                        "parse_version": 1,
                        "section_id": "section-7",
                        "section_revision": 1,
                        "locator": "paragraph:7",
                        "page": None,
                        "coordinate_system": "ooxml",
                        "quote": "供应商必须在本市注册。",
                        "quote_sha256": "c" * 64,
                    }
                ],
            }
        ],
    }


def test_task0010_word_and_pdf_are_real_parseable_chinese_reports() -> None:
    report = _report()
    docx_bytes = render_docx(report)  # type: ignore[arg-type]
    pdf_bytes = render_pdf(report)  # type: ignore[arg-type]
    assert docx_bytes.startswith(b"PK")
    assert pdf_bytes.startswith(b"%PDF-")
    with ZipFile(BytesIO(docx_bytes)) as archive:
        assert "word/document.xml" in archive.namelist()
        xml = archive.read("word/document.xml").decode("utf-8")
        assert "注册地要求需复核" in xml
        assert "doc-old-version" in xml
        assert "suspected" in xml
        assert "confirmed" not in xml
    parsed_docx = Document(BytesIO(docx_bytes))
    assert any(
        "中文项目" in cell.text
        for table in parsed_docx.tables
        for row in table.rows
        for cell in row.cells
    )
    parsed_pdf = PdfReader(BytesIO(pdf_bytes))
    assert len(parsed_pdf.pages) >= 1
    pdf_text = "\n".join(page.extract_text() or "" for page in parsed_pdf.pages)
    assert "招标文件合规审查报告" in pdf_text
    assert "suspected" in pdf_text
    assert "confirmed" not in pdf_text


def test_task0022_zero_candidate_reports_disclose_actual_run_limits() -> None:
    report = _report()
    report.update(
        {
            "schema_version": "task0010-report.v2",
            "template_version": "task0010-v2",
            "findings": [],
            "assessment_scope": {
                "candidate_discovery_status": "completed_with_incomplete_skill",
                "candidate_count": 0,
                "skill_coverage": [
                    {
                        "skill_id": "qualification_review",
                        "status": "completed_no_candidates",
                        "candidate_count": 0,
                        "evaluated_rule_count": 1,
                        "reason": None,
                    },
                    {
                        "skill_id": "technical_review",
                        "status": "completed_no_candidates",
                        "candidate_count": 0,
                        "evaluated_rule_count": 1,
                        "reason": None,
                    },
                    {
                        "skill_id": "evaluation_review",
                        "status": "incomplete",
                        "candidate_count": 0,
                        "evaluated_rule_count": 1,
                        "reason": "ambiguous_score_table",
                    },
                    {
                        "skill_id": "procedure_review",
                        "status": "incomplete",
                        "candidate_count": 0,
                        "evaluated_rule_count": 1,
                        "reason": "ambiguous_tender_times",
                    },
                    {
                        "skill_id": "contract_review",
                        "status": "not_applicable",
                        "candidate_count": 0,
                        "evaluated_rule_count": 1,
                        "reason": None,
                    },
                ],
                "legal_framework_confirmed": False,
                "retrieval_status_counts": {"context_incomplete": 5},
                "retrieval_context_errors": {"CONTEXT_INVALID": 5},
                "retrieval_hit_count": 0,
                "citable_legal_evidence_count": 0,
                "retrieval_algorithms": ["bm25-char-cosine-rrf-5"],
                "model_provider": "offline-verification",
                "model_name": "synthetic-protocol-adapter",
                "no_findings_confirmation": "confirm_no_findings",
                "source_artifact_hashes": {
                    "candidate_discovery": "d" * 64,
                    "legal_retrieval": "e" * 64,
                    "model_analysis": "f" * 64,
                    "project_context": "a" * 64,
                },
            },
        }
    )
    schema = json.loads(
        (Path(__file__).resolve().parents[1] / "contracts/schemas/review-report.schema.json")
        .read_text(encoding="utf-8")
    )
    Draft202012Validator(schema).validate(report)

    parsed_word = Document(BytesIO(render_docx(report)))
    word_text = "\n".join(paragraph.text for paragraph in parsed_word.paragraphs)
    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(render_pdf(report))).pages
    )
    expected_parts = (
        "评审办法",
        "未完成判断",
        "评分表存在歧义",
        "采购程序",
        "采购时间存在歧义",
        "0候选仅表示已运行规则范围内未产生候选",
        "法律适用框架尚未确认",
        "CONTEXT_INVALID",
        "未取得可确认、可引用的法规依据",
        "OfflineReviewGateway",
        "本次未调用真实LLM",
        "confirm_no_findings",
        "不构成专业法律审查确认",
        "全部合规",
    )
    for output_index, output in enumerate((word_text, pdf_text)):
        normalized = "".join(output.split())
        missing = [part for part in expected_parts if part not in normalized]
        assert not missing, f"output {output_index} omitted required scope statements: {missing}"


@pytest.mark.parametrize(
    ("action", "expected_status"),
    (("accept", "accepted"), ("modify", "modified"), ("reject", "rejected")),
)
def test_task0010_human_projection_keeps_evidence_status_separate(
    action: str, expected_status: str
) -> None:
    principal = Principal("tenant-test", "reviewer-test", "reviewer", Role.REVIEWER, "session")
    source = {
        "finding_id": "finding-test",
        "evidence_status": "insufficient_evidence",
        "review_status": "pending",
        "title": "候选风险",
        "evidence": [],
    }
    modifications = {"title": "修改后的候选风险"} if action == "modify" else None
    projected = _decision_projection(
        source,
        action,
        modifications,
        principal,
        "已阅读并接受记录。",
        datetime.fromisoformat("2026-10-14T00:00:00+00:00"),
    )
    assert projected["review_status"] == expected_status
    assert projected["evidence_status"] == "insufficient_evidence"
    assert projected["human_review"]["reason"] == "已阅读并接受记录。"


def _reference(article: str, title: str, reason: str) -> dict:
    return {
        "evidence_id": "d" * 64,
        "kind": "law",
        "title": title,
        "article": article,
        "quote": "约定的违约金过分高于造成的损失的，可以根据当事人的请求予以适当减少。",
        "source_uri": "https://www.gov.cn/example/article",
        "source_hash": "e" * 64,
        "locator": f"article:{article}",
        "version_status": "unverified",
        "confirmation_reason": reason,
        "applicable": False,
    }


def test_reference_only_clauses_render_in_word_and_pdf_without_becoming_evidence() -> None:
    """Unconfirmed clauses must be visible as references and must not raise any gate."""
    report = _report()
    finding = report["findings"][0]
    finding["evidence"] = []
    finding["evidence_status"] = "insufficient_evidence"
    finding["sources"] = ["model"]
    finding["legal_references"] = [
        _reference("第五百八十五条", "中华人民共和国民法典", "unverified"),
        _reference("第二十条", "中华人民共和国政府采购法实施条例", "applicability_unverified"),
    ]

    docx_bytes = render_docx(report)  # type: ignore[arg-type]
    pdf_bytes = render_pdf(report)  # type: ignore[arg-type]
    word_text = "".join(
        paragraph.text for paragraph in Document(BytesIO(docx_bytes)).paragraphs
    )
    pdf_text = "".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_bytes)).pages)

    for label, text in (("docx", word_text), ("pdf", pdf_text)):
        compact = "".join(text.split())
        assert "可能相关依据" in compact, label
        assert "不计入已核验证据" in compact, label
        assert "第五百八十五条" in compact, label
        assert "中华人民共和国民法典" in compact, label
        assert "第二十条" in compact, label
        assert "unverified" in compact, label
        assert "applicability_unverified" in compact, label
        # The reference section must never read as a violation finding.
        tail = compact.split("可能相关依据", 1)[1]
        assert "违反" not in tail, label
        assert "违法依据" not in tail, label

    # Finding-level isolation: references never enter the evidence or source collections.
    assert finding["evidence"] == []
    assert finding["sources"] == ["model"]
    assert finding["evidence_status"] == "insufficient_evidence"


def _v3_scope(product_execution: dict[str, object]) -> dict[str, object]:
    return {
        "candidate_discovery_status": "completed",
        "candidate_count": 1,
        "skill_coverage": [],
        "legal_framework_confirmed": False,
        "retrieval_status_counts": {"completed": 1},
        "retrieval_context_errors": {},
        "retrieval_hit_count": 1,
        "citable_legal_evidence_count": 0,
        "retrieval_algorithms": ["bm25-char-cosine-rrf-5"],
        "model_provider": "bailian",
        "model_name": "qwen3.8-flash",
        "no_findings_confirmation": None,
        "source_artifact_hashes": {
            "candidate_discovery": "d" * 64,
            "legal_retrieval": "e" * 64,
            "model_analysis": "f" * 64,
            "project_context": "a" * 64,
        },
        "product_execution": product_execution,
    }


def test_product_execution_projection_marks_saturation_and_unknown_calls_partial() -> None:
    partial = project_product_execution(
        {
            "status": "partial",
            "saturated_batches": 1,
            "partial_batches": 1,
            "section_count": 12,
        },
        unknown_model_call_count=1,
        batches=[
            {"status": "completed", "section_count": 10},
            {"status": "saturated", "section_count": 2},
        ],
    )
    assert partial["completion_status"] == "partial"
    assert partial["coverage_status"] == "partial"
    assert partial["unprocessed_units"] == 2
    assert partial["saturated"] is True
    assert partial["truncated"] is True
    assert partial["unknown_model_call_count"] == 1
    assert {"discovery_saturated", "model_call_outcome_unknown"} <= set(partial["reason_codes"])

    complete = project_product_execution(
        {"status": "complete", "saturated_batches": 0, "partial_batches": 0},
        unknown_model_call_count=0,
        batches=[{"status": "completed", "section_count": 5}],
    )
    assert complete["completion_status"] == "complete"
    assert complete["coverage_status"] == "complete"
    assert complete["unprocessed_units"] == 0


def test_v3_report_discloses_partial_coverage_in_word_and_pdf() -> None:
    report = _report()
    report["schema_version"] = "task0010-report.v3"
    report["template_version"] = "task0010-v3"
    report["run_id"] = "00000000-0000-4000-8000-000000000001"
    report["result_id"] = "00000000-0000-4000-8000-000000000002"
    report["report_id"] = "00000000-0000-4000-8000-000000000003"
    report["assessment_scope"] = _v3_scope(
        {
            "completion_status": "partial",
            "coverage_status": "partial",
            "unprocessed_units": 2,
            "truncated": True,
            "saturated": True,
            "reason_codes": ["discovery_saturated"],
            "unknown_model_call_count": 0,
        }
    )
    documents, registry, errors = load_contracts(Path(__file__).resolve().parents[1])
    api, api_errors = load_openapi(Path(__file__).resolve().parents[1])
    assert not errors + api_errors
    validator = validator_for_contract(
        Path(__file__).resolve().parents[1],
        "contracts/schemas/review-report.schema.json",
        documents,
        registry,
        api,
    )
    validator.validate(report)

    word_text = "".join(
        paragraph.text for paragraph in Document(BytesIO(render_docx(report))).paragraphs  # type: ignore[arg-type]
    )
    pdf_text = "".join(
        page.extract_text() or ""
        for page in PdfReader(BytesIO(render_pdf(report))).pages  # type: ignore[arg-type]
    )
    for label, text in (("docx", word_text), ("pdf", pdf_text)):
        compact = "".join(text.split())
        assert "部分完成" in compact, label
        assert "不得解释为全部内容已审查完成" in compact, label
        assert "未处理单元" in compact, label
        assert "饱和" in compact, label


def test_v3_report_discloses_project_adopted_legal_references() -> None:
    report = _report()
    report["schema_version"] = "task0010-report.v3"
    report["template_version"] = "task0010-v3"
    report["run_id"] = "00000000-0000-4000-8000-000000000011"
    report["result_id"] = "00000000-0000-4000-8000-000000000012"
    report["report_id"] = "00000000-0000-4000-8000-000000000013"
    report["assessment_scope"] = _v3_scope(
        {
            "completion_status": "complete",
            "coverage_status": "complete",
            "unprocessed_units": 0,
            "truncated": False,
            "saturated": False,
            "reason_codes": ["discovery_complete"],
            "unknown_model_call_count": 0,
        }
    )
    report["assessment_scope"]["legal_adoption_policy"] = {
        "policy_id": "ADOPT-2026-10-18",
        "decision_ref": "DEC-009",
        "adopted_on": "2026-10-18",
        "independent_legal_verification": False,
        "version_count": 43,
        "fragment_count": 3120,
        "corpus_index_sha256": "a" * 64,
        "manifest_sha256": "b" * 64,
        "statement": "项目采用库内法规版本，未做独立法律核验。",
    }
    report["assessment_scope"]["adopted_legal_reference_count"] = 1
    report["assessment_scope"]["configured_rule_law_binding_count"] = 1
    report["findings"][0]["legal_references"] = [
        {
            "evidence_id": "e" * 64,
            "kind": "law",
            "title": "中华人民共和国招标投标法实施条例",
            "article": "第三十二条",
            "quote": "招标人不得以不合理的条件限制、排斥潜在投标人或者投标人。",
            "source_uri": "artifact://legal-text/" + "f" * 64,
            "source_hash": "f" * 64,
            "version_status": "unverified",
            "confirmation_reason": "outside_verification_horizon",
            "adoption_status": "project_adopted",
            "selected_by_model": True,
            "rule_law_binding_id": "BIND-QUAL-LOCAL-32",
        }
    ]
    documents, registry, errors = load_contracts(Path(__file__).resolve().parents[1])
    api, api_errors = load_openapi(Path(__file__).resolve().parents[1])
    assert not errors + api_errors
    validator = validator_for_contract(
        Path(__file__).resolve().parents[1],
        "contracts/schemas/review-report.schema.json",
        documents,
        registry,
        api,
    )
    validator.validate(report)

    word_text = "".join(
        paragraph.text for paragraph in Document(BytesIO(render_docx(report))).paragraphs  # type: ignore[arg-type]
    )
    pdf_text = "".join(
        page.extract_text() or ""
        for page in PdfReader(BytesIO(render_pdf(report))).pages  # type: ignore[arg-type]
    )
    for label, text in (("docx", word_text), ("pdf", pdf_text)):
        compact = "".join(text.split())
        assert "法规采用政策" in compact, label
        assert "ADOPT-2026-10-18" in compact, label
        assert "项目采用法规库参考" in compact, label
        assert "项目采用库内版本" in compact, label
        assert "BIND-QUAL-LOCAL-32" in compact, label
        assert "未做独立法律核验" in compact, label
